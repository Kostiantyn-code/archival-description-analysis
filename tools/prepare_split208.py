"""Prepare the published Split inventory as an auditable input matrix.

One Python command writes the analysis XLSX, a source-row mapping and a
conversion manifest. No classification or translation occurs in this step.
"""

from __future__ import annotations

import argparse
from collections import Counter
from datetime import date
import hashlib
import json
from pathlib import Path
import re
import xml.etree.ElementTree as ET
import zipfile

VERSION = "split208-import-1.1"
SOURCE_URL = "https://www.das.hr/wp-content/uploads/2021/09/HR-DAST-208-GRADSKA-KOMISIJA-ZA-RATNU-%C5%A0TETU-SPLIT_serija-Ratna-%C5%A1teta-MNO-Split.xlsx"
LANDING_URL = "https://www.das.hr/obavijesna-pomagala/"
EXPECTED_SHA256 = "3bba78d372ffdd33c81d975e0ee908e28dbe98fb84ae781c2e37c4c5cb557785"
NS = {"s": "http://schemas.openxmlformats.org/spreadsheetml/2006/main"}
FIELDS = [
    "claim_number",
    "local_committee",
    "applicant",
    "register_number",
    "registration_date",
    "damage_description",
    "technical_unit",
]


def read_rows(path: Path) -> list[dict]:
    with zipfile.ZipFile(path) as archive:
        shared = []
        if "xl/sharedStrings.xml" in archive.namelist():
            shared = [
                "".join(item.itertext())
                for item in ET.fromstring(archive.read("xl/sharedStrings.xml"))
            ]
        rows = []
        for row in ET.fromstring(archive.read("xl/worksheets/sheet1.xml")).findall(
            "s:sheetData/s:row", NS
        ):
            source_row = int(row.attrib["r"])
            if source_row < 3:
                continue
            cells = {}
            for cell in row.findall("s:c", NS):
                raw = cell.find("s:v", NS)
                value = "" if raw is None else raw.text or ""
                if cell.get("t") == "s" and value:
                    value = shared[int(value)]
                elif cell.get("t") == "inlineStr":
                    inline = cell.find("s:is", NS)
                    value = "" if inline is None else "".join(inline.itertext())
                cells[re.sub(r"\d", "", cell.attrib["r"])] = value
            values = [cells.get(column, "") for column in "ABCDEFG"]
            if any(value.strip() for value in values):
                rows.append({"source_row": source_row, **dict(zip(FIELDS, values))})
        return rows


def date_status(value: str) -> str:
    if not value.strip():
        return "missing"
    match = re.fullmatch(
        r"(\d{1,2})\.([IVX]+|\d{1,2})\.(\d{4})",
        re.sub(r"\s+", "", value.upper()).rstrip("."),
    )
    if not match:
        return "unparsed"
    romans = {
        v: i for i, v in enumerate("I II III IV V VI VII VIII IX X XI XII".split(), 1)
    }
    day, month, year = match.groups()
    try:
        date(int(year), int(month) if month.isdigit() else romans[month], int(day))
    except (ValueError, KeyError):
        return "invalid"
    return "parsed"


def prepare(source: Path, output: Path) -> dict:
    digest = hashlib.sha256(source.read_bytes()).hexdigest()
    if digest != EXPECTED_SHA256:
        raise ValueError(
            "Source SHA-256 differs from the documented inventory snapshot. Review the new source before adapting the converter."
        )
    rows = read_rows(source)
    matrix = [
        ["Archive", "Državni arhiv u Splitu", None, "Language", "hr"] + [None] * 7,
        ["Fonds", "HR-DAST-208", None, "Source", SOURCE_URL] + [None] * 7,
        ["Inventory", "Ratna šteta – MNO Split", None, "SHA-256", digest] + [None] * 7,
        [
            "No.",
            "Title",
            "Dates",
            "",
            "Notes",
            "Source row",
            "Claim number",
            "Local committee",
            "Applicant",
            "Register number",
            "Technical unit",
            "Date status",
        ],
    ]
    mapping = []
    for row in rows:
        source_row = row["source_row"]
        record_id = f"HR-DAST-208_MNO-Split_r{source_row:05d}"
        status = date_status(row["registration_date"])
        notes = json.dumps(
            {
                "source_row": source_row,
                "date_basis": "registration_date",
                "date_status": status,
                **{
                    key: row[key]
                    for key in FIELDS
                    if key not in {"damage_description", "registration_date"}
                },
            },
            ensure_ascii=False,
        )
        matrix.append(
            [
                record_id,
                row["damage_description"],
                row["registration_date"],
                None,
                notes,
                source_row,
                row["claim_number"],
                row["local_committee"],
                row["applicant"],
                row["register_number"],
                row["technical_unit"],
                status,
            ]
        )
        mapping.append(
            {
                "record_id": record_id,
                "source_row": source_row,
                "output_row": len(matrix),
            }
        )
    manifest = {
        "algorithm_version": VERSION,
        "source_url": SOURCE_URL,
        "landing_url": LANDING_URL,
        "source_sha256": digest,
        "source_member": "xl/worksheets/sheet1.xml",
        "source_rows": [rows[0]["source_row"], rows[-1]["source_row"]],
        "records": len(rows),
        "nonempty_descriptions": sum(
            bool(row["damage_description"].strip()) for row in rows
        ),
        "date_status_counts": dict(
            Counter(date_status(row["registration_date"]) for row in rows)
        ),
        "original_claim_numbers_repeated": {
            key: value
            for key, value in Counter(
                row["claim_number"] for row in rows if row["claim_number"].strip()
            ).items()
            if value > 1
        },
        "mapping": {
            "A": "derived unique identifier, based on source row",
            "B": "source F, verbatim damage description",
            "C": "source E, verbatim registration date (not the date of damage)",
            "D": "blank: no folio counts in source",
            "E": "JSON provenance and source A–D/G values; date quality flag",
            "F": "source row number",
            "G": "source A",
            "H": "source B",
            "I": "source C",
            "J": "source D",
            "K": "source G",
            "L": "derived date quality flag",
        },
        "rules": [
            "Keep source order and every nonempty row from row 3 onward.",
            "No translation, spelling correction, lowercasing, category assignment, deduplication or date imputation.",
            "Retain blank descriptions and repeated claim numbers. The technical identifier is not an archival file number.",
            "Only the published Ratna šteta – MNO Split inventory is represented, not the whole fonds.",
            "The program reads A–E. Additional columns preserve accessible source metadata.",
        ],
    }
    output.mkdir(parents=True, exist_ok=True)
    output_file = output / "HR-DAST-208_MNO-Split_input.xlsx"
    write_xlsx(matrix, output_file)
    manifest.update(
        output_file=output_file.name,
        output_sha256=hashlib.sha256(output_file.read_bytes()).hexdigest(),
        export_engine="openpyxl",
        output_rows_including_metadata=len(matrix),
    )
    for name, data in [
        ("split208_row_mapping.json", mapping),
        ("split208_conversion_manifest.json", manifest),
    ]:
        (output / name).write_text(
            json.dumps(data, ensure_ascii=False, indent=2) + "\n", encoding="utf-8"
        )
    return manifest


def write_xlsx(matrix: list[list], destination: Path) -> None:
    """Preserve source strings as text, including a literal leading equals sign."""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "MNO Split"
    sheet.sheet_view.showGridLines = False
    widths = [43, 77, 20, 4, 92, 12, 20, 24, 34, 20, 18, 15]
    for row_number, values in enumerate(matrix, 1):
        for column, value in enumerate(values, 1):
            cell = sheet.cell(row_number, column, value)
            if isinstance(value, str):
                cell.data_type = "s"
                cell.number_format = "@"
            cell.font = Font(name="Arial", size=10, color="172F43")
            cell.alignment = Alignment(vertical="top", wrap_text=column in {2, 3, 5})
            if row_number == 4:
                cell.font = Font(name="Arial", size=10, color="FFFFFF", bold=True)
                cell.fill = PatternFill("solid", fgColor="26465F")
        sheet.row_dimensions[row_number].height = 32 if row_number == 4 else 36
    for column, width in enumerate(widths, 1):
        sheet.column_dimensions[get_column_letter(column)].width = width
    for row in range(1, 4):
        sheet.cell(row, 1).font = Font(name="Arial", size=10, bold=True)
    sheet.freeze_panes = "B5"
    workbook.save(destination)
    workbook.close()


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("source", type=Path)
    parser.add_argument("output", type=Path)
    args = parser.parse_args()
    result = prepare(args.source, args.output)
    print(
        json.dumps(
            {
                key: result[key]
                for key in (
                    "algorithm_version",
                    "records",
                    "nonempty_descriptions",
                    "date_status_counts",
                    "source_sha256",
                )
            },
            indent=2,
        )
    )
