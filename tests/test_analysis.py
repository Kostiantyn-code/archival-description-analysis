import importlib.util
import sys
import tempfile
import unittest
from pathlib import Path
import yaml
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))
import fond_analysis as f


class AnalysisTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        f.configure_analysis(yaml)
        f.configure_sources(yaml)
        cls.uk, cls.amb, _, cls.mac, _ = f.load_dictionaries(yaml)
        ru, amb, *_ = f.load_dictionaries(yaml, 'ru')
        f.LANGUAGE_CATEGORIES.update(uk=cls.uk, ru=ru)
        f.LANGUAGE_AMBIGUITIES.update(uk=cls.amb, ru=amb)

    def topics(self, text, lang='ru'):
        return f.classify_title(text, self.uk, self.amb, list(self.mac), lang)[0]

    def test_status(self):
        for title in ['В И Б У Л А', 'В\u00a0И\u00a0Б\u00a0У\u00a0Л\u00a0И', 'ВИБУЛО', 'В Ы Б Ы Л А']:
            self.assertEqual(f.detect_status('33-49', title), 'withdrawn')
        self.assertEqual(f.detect_status('1', 'Справа про вибуття солдатів'), 'case')
        self.assertEqual(f.detect_status('476', 'Недействующая опись'), 'inventory')
        for title in ['1850', '1850 год', '1896, 1899 – 1901 рік']:
            self.assertEqual(f.detect_status('', title), 'year_heading')
        self.assertEqual(f.detect_status('', 'Загальний розділ'), 'group_heading')
        self.assertEqual(f.detect_status('', '', '1850'), 'unknown')

    def test_language(self):
        self.assertEqual(f.detect_title_language('Дело о розыске разных лиц'), 'ru')
        self.assertEqual(f.detect_title_language('Справа про відкриття школи та навчання дітей'), 'uk')
        self.assertEqual(f.detect_title_language('Николаев 1850'), 'undetermined')

    def test_historical_examples(self):
        for title, topic in [
            ('Указы Государственной военной коллегии', 'military_affairs'),
            ('Дело о нарушении карантинного режима', 'healthcare'),
            ('Дело о нарушении карантинного режима', 'law_and_police'),
            ('Открытие Севастопольского порта для торговых судов', 'economy'),
            ('Дело о политической проверке разных лиц', 'politics'),
            ('Дело о розыске разных лиц', 'law_and_police'),
            ('Дело о мощении улиц', 'urban_infrastructure'),
            ('Открытие народного училища и обучение детей', 'education'),
            ('Списки книг, запрещенных Комитетом цензуры иностранной', 'culture'),
            ('Дело о распространении раскола', 'religion'),
            ('Дело о материальной помощи вдове', 'population_and_society'),
            ('Дело о назначении на должность', 'public_administration')
            ,('Дело о предупреждении еврейского погрома', 'politics')
            ,('О недопущении сверхсметных издержек', 'economy')
            ,('Дело о запрещении фотографических снимков', 'culture')
        ]:
            with self.subTest(title=title, topic=topic):
                self.assertIn(topic, self.topics(title))

    def test_context_not_subject(self):
        for title, forbidden in [
            ('Дело об аресте купца Иванова', 'economy'),
            ('Дело о розыске лекаря Иванова', 'healthcare'),
            ('Дело об аресте отставного капитана', 'military_affairs'),
            ('Переписка с военным губернатором', 'military_affairs'),
            ('Журнал входящих секретных документов', 'culture'),
            ('Духовное завещание купца', 'religion'),
            ('Дело об освещении улиц', 'education'),
            ('Дело об аресте бывшего студента университета', 'education'),
            ('Надзор полиции за купцом', 'politics')
            ,('Розыск крестьянки раскольницы', 'religion')
            ,('Дело о розыске скопца Купцова', 'religion')
            ,('Дело об аресте рабочего судостроительного завода', 'economy')
        ]:
            with self.subTest(title=title):self.assertNotIn(forbidden, self.topics(title))

    def test_healthcare_requires_subject_not_incidental_institution(self):
        for title, language in [
            ('Справа про нагородження чиновника Севастопольського карантину', 'uk'),
            ('Листування про квартиру наглядача військового шпиталю', 'uk'),
            ('Дело о розыске лекаря Иванова', 'ru'),
            ('Дело о награждении чиновника карантина', 'ru'),
        ]:
            with self.subTest(title=title):
                self.assertNotIn('healthcare', self.topics(title, language))

        for title, language in [
            ('Справа про ремонт міської лікарні', 'uk'),
            ('Справа про призначення міського лікаря на посаду', 'uk'),
            ('Заходи боротьби з епідемією холери', 'uk'),
            ('Дело о строительстве больницы', 'ru'),
            ('Дело о карантинных правилах', 'ru'),
        ]:
            with self.subTest(title=title):
                self.assertIn('healthcare', self.topics(title, language))

    def test_cache_is_language_specific(self):
        before = self.topics('Дело об аресте купца')
        self.topics('Справа про арешт купця', 'uk')
        self.assertEqual(before, self.topics('Дело об аресте купца'))
        self.assertEqual(f.MATCH_LANGUAGE.get(), 'uk')

    def test_years(self):
        issues=[]
        self.assertEqual(f.parse_years('24 січня\n24 червня 1850', 3, '1', issues, 'ЦДІАК'), (1850,1850))
        f.parse_years('1860–1859',4,'2',issues,'ЦДІАК')
        self.assertTrue(any(i.level=='ERROR' for i in issues))
        self.assertEqual(f.parse_years('18648',5,'3',[], 'ЦДІАК'),(None,None))

    def test_identifiers(self):
        self.assertEqual(f.normalize_case_id('2712 а'),'2712а')
        self.assertEqual(f.normalize_case_id('33 – 49'),'33-49')

    def test_dictionary_completeness(self):
        self.assertEqual({c.id for c in f.LANGUAGE_CATEGORIES['uk']},
                         {c.id for c in f.LANGUAGE_CATEGORIES['ru']})
        self.assertEqual(len(f.LANGUAGE_CATEGORIES['ru']),11)
        for language, categories in f.LANGUAGE_CATEGORIES.items():
            for c in categories:
                self.assertGreaterEqual(c.minimum_score, 1)
                self.assertTrue(c.strong_terms or c.strong_phrases or c.context_rules)

    def test_healthcare_is_configured_bilingually(self):
        healthcare_uk = next(c for c in f.LANGUAGE_CATEGORIES['uk'] if c.id == 'healthcare')
        healthcare_ru = next(c for c in f.LANGUAGE_CATEGORIES['ru'] if c.id == 'healthcare')
        self.assertEqual(healthcare_uk.export_filename, 'охорона_здоров’я.xlsx')
        self.assertEqual(healthcare_ru.export_filename, 'охорона_здоров’я.xlsx')
        self.assertIn('карантин', healthcare_uk.context_only_terms)
        self.assertIn('карантин', healthcare_ru.context_only_terms)

    def test_classification_type_is_multilabel_aware(self):
        record = f.Record(2, 'Опис 1', '1', '1', 'Заголовок', '', '', '', 'case', 1850)
        self.assertEqual(f.classification_type(record), 'unclassified')
        record.context_categories = ['healthcare']
        self.assertEqual(f.classification_type(record), 'context_only')
        record.categories = ['healthcare']
        self.assertEqual(f.classification_type(record), 'subject_single')
        record.categories.append('law_and_police')
        self.assertEqual(f.classification_type(record), 'subject_multilabel')

    def test_runtime_category_list_controls_export_files(self):
        old_input = f.INPUT_FILE
        old_thematic = f.THEMATIC_DIR
        old_config = f.ANALYSIS_CONFIG
        try:
            with tempfile.TemporaryDirectory() as directory:
                directory = Path(directory)
                workbook = Workbook()
                workbook.remove(workbook.active)
                for name in f.SHEET_NAMES:
                    sheet = workbook.create_sheet(name)
                    sheet.append(['№ з/п', 'Заголовок справи', 'Крайні дати', 'Аркуші', 'Примітки'])
                source_sheet = workbook['Опис 1']
                header = source_sheet['A1']
                header.font = Font(name='Arial', size=11, bold=True)
                header.fill = PatternFill('solid', fgColor='E5EDF4')
                header.alignment = Alignment(wrap_text=True)
                header.border = Border(bottom=Side(style='thin', color='808080'))
                source_sheet.row_dimensions[1].height = 26.4
                source_sheet.column_dimensions['A'].width = 12.5
                source_sheet.sheet_view.zoomScale = 90
                source_sheet.freeze_panes = 'B2'
                source_sheet.append(['1', 'Про лікарню', '1850', '2', ''])
                source_sheet.append(['2', 'Про театр', '1851', '3', ''])
                source_sheet['B2'].font = Font(name='Arial', italic=True)
                input_path = directory / 'input.xlsx'
                workbook.save(input_path)

                f.INPUT_FILE = input_path
                f.THEMATIC_DIR = directory / 'exports'
                f.ANALYSIS_CONFIG = {
                    'exports': {'enabled': True, 'source_columns': 5,
                                'xlsx': True, 'csv': False}
                }
                record = f.Record(
                    2, 'Опис 1', '1', '1', 'Про лікарню', '1850', '2', '',
                    'case', 1850, categories=['healthcare'],
                    category_labels=['Охорона здоров’я і санітарія'],
                    scores={'healthcare': 4},
                )
                category = next(c for c in self.uk if c.id == 'healthcare')
                f.export_thematic_selections(__import__('openpyxl'), [record], [category])
                self.assertTrue((f.THEMATIC_DIR / 'охорона_здоров’я.xlsx').exists())
                self.assertFalse((f.THEMATIC_DIR / 'культура.xlsx').exists())
                self.assertTrue((f.THEMATIC_DIR / 'некласифіковані.xlsx').exists())
                self.assertTrue((f.THEMATIC_DIR / 'службові_записи.xlsx').exists())
                from openpyxl import load_workbook
                result = load_workbook(
                    f.THEMATIC_DIR / 'охорона_здоров’я.xlsx', read_only=False
                )
                self.assertEqual(result.sheetnames, list(f.SHEET_NAMES))
                self.assertEqual(
                    sum(max(0, sum(1 for _ in ws.iter_rows()) - 1)
                        for ws in result.worksheets),
                    1,
                )
                exported = result['Опис 1']
                self.assertEqual(exported.max_row, 2)
                self.assertEqual(exported['A2'].value, '1')
                self.assertEqual(exported['B2'].value, 'Про лікарню')
                self.assertTrue(exported['A1'].font.bold)
                self.assertEqual(exported['A1'].font.name, 'Arial')
                self.assertEqual(exported['A1'].fill.fill_type, 'solid')
                self.assertEqual(exported['A1'].fill.fgColor.rgb, '00E5EDF4')
                self.assertTrue(exported['A1'].alignment.wrap_text)
                self.assertEqual(exported['A1'].border.bottom.style, 'thin')
                self.assertAlmostEqual(exported.row_dimensions[1].height, 26.4)
                self.assertAlmostEqual(exported.column_dimensions['A'].width, 12.5)
                self.assertEqual(exported.sheet_view.zoomScale, 90)
                self.assertEqual(exported.freeze_panes, 'B2')
                self.assertTrue(exported['B2'].font.italic)
                result.close()
        finally:
            f.INPUT_FILE = old_input
            f.THEMATIC_DIR = old_thematic
            f.ANALYSIS_CONFIG = old_config

    def test_disabled_category_is_not_loaded(self):
        old_config_dir = f.CONFIG_DIR
        try:
            with tempfile.TemporaryDirectory() as directory:
                directory = Path(directory)
                config = {
                    'dictionary_version': 'test',
                    'classification': {'minimum_score': 3},
                    'macroblocks': {'cultural': {'label': 'Культурне життя'}},
                    'categories': [
                        {
                            'id': 'culture', 'label': 'Культура',
                            'macroblock': 'cultural', 'enabled': True,
                            'export': 'культура.xlsx',
                            'dictionaries': {
                                'uk': str(ROOT / 'dictionaries/cultural/culture.yaml')
                            },
                        },
                        {
                            'id': 'education', 'label': 'Освіта',
                            'macroblock': 'cultural', 'enabled': False,
                            'export': 'освіта.xlsx',
                            'dictionaries': {
                                'uk': str(ROOT / 'dictionaries/cultural/education.yaml')
                            },
                        },
                    ],
                }
                (directory / 'categories.yaml').write_text(
                    yaml.safe_dump(config, allow_unicode=True), encoding='utf-8'
                )
                f.CONFIG_DIR = directory
                categories, *_ = f.load_dictionaries(yaml, 'uk')
                self.assertEqual([category.id for category in categories], ['culture'])
        finally:
            f.CONFIG_DIR = old_config_dir

    def test_source_config(self):
        self.assertEqual(f.SOURCE_CONFIG['ЦДІАК']['language'],'ru')
        self.assertEqual(f.SOURCE_CONFIG['ЦДІАК']['fond'],'356')
        self.assertEqual(f.SOURCE_CONFIG['Опис 1']['fond'],'230')
        self.assertEqual(f.SOURCE_CONFIG['ЦДІАК']['inventory'],'')

    def test_chronology(self):
        r = f.Record(2,'Опис 1','1','1','Заголовок','','','','case',1850)
        self.assertEqual(f.chronology_basis(r),'section_heading')
        self.assertEqual(f.get_decade(r),1850)
        r.start_year=1865;r.end_year=1858
        self.assertEqual(f.chronology_basis(r),'invalid_dates')
        self.assertIsNone(f.get_decade(r))
        r.start_year=1850;r.end_year=1855
        self.assertEqual(f.chronology_basis(r),'dates_field')


if __name__ == '__main__':
    unittest.main(verbosity=2)
