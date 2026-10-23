"""The public-source converter preserves cells, mappings and literal text."""

import hashlib
import json
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

import openpyxl
from tools import prepare_split208 as converter


class ConverterTests(unittest.TestCase):
    def test_source_mapping_literal_text_and_date_flags(self):
        with tempfile.TemporaryDirectory() as temporary:
            folder = Path(temporary)
            source = folder / "source.xlsx"
            book = openpyxl.Workbook()
            sheet = book.active
            sheet.append(["Archive title"])
            sheet.append(
                [
                    "claim",
                    "committee",
                    "applicant",
                    "register",
                    "date",
                    "damage",
                    "unit",
                ]
            )
            sheet.append(["7", "Split", "Name", "12", "31.II.1945.", "  roba  ", "1"])
            sheet.append(["7", "Split", "Name", "13", "", "=literal text", "2"])
            sheet["F4"].data_type = "s"
            sheet.append(["8", "Split", "Name", "14", "1.III.1945.", None, "3"])
            book.save(source)
            book.close()
            digest = hashlib.sha256(source.read_bytes()).hexdigest()
            with patch.object(converter, "EXPECTED_SHA256", digest):
                manifest = converter.prepare(source, folder / "output")
            output = folder / "output" / manifest["output_file"]
            result = openpyxl.load_workbook(output, data_only=False)
            try:
                sheet = result.active
                self.assertEqual(sheet["B5"].value, "  roba  ")
                self.assertEqual(sheet["B6"].value, "=literal text")
                self.assertEqual(sheet["B6"].data_type, "s")
                self.assertEqual(sheet["D4"].value, None)
                self.assertEqual(sheet["E1"].value, "hr")
                self.assertEqual(sheet["A5"].value, "HR-DAST-208_MNO-Split_r00003")
                self.assertEqual(sheet["L5"].value, "invalid")
                self.assertEqual(sheet["L6"].value, "missing")
                self.assertEqual(sheet["L7"].value, "parsed")
                self.assertEqual(sheet.max_row, 7)
            finally:
                result.close()
            self.assertEqual(manifest["records"], 3)
            self.assertEqual(manifest["nonempty_descriptions"], 2)
            self.assertEqual(manifest["original_claim_numbers_repeated"], {"7": 2})
            self.assertEqual(
                manifest["output_sha256"],
                hashlib.sha256(output.read_bytes()).hexdigest(),
            )
            mapping = json.loads(
                (folder / "output/split208_row_mapping.json").read_text()
            )
            self.assertEqual([row["source_row"] for row in mapping], [3, 4, 5])
            self.assertEqual([row["output_row"] for row in mapping], [5, 6, 7])

    def test_changed_source_is_rejected_before_conversion(self):
        with tempfile.TemporaryDirectory() as temporary:
            source = Path(temporary) / "changed.xlsx"
            source.write_bytes(b"changed source")
            with self.assertRaisesRegex(ValueError, "SHA-256"):
                converter.prepare(source, Path(temporary) / "output")
            self.assertFalse((Path(temporary) / "output").exists())


if __name__ == "__main__":
    unittest.main()
