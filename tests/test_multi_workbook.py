"""New workbook metadata and independent batch results."""
from __future__ import annotations

import contextlib
import csv
import io
import json
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

import openpyxl
import yaml

import sys
sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
import fond_analysis as f


def workbook(path: Path, descriptions) -> None:
    book = openpyxl.Workbook()
    book.remove(book.active)
    for sheet_name, archive, fond, inventory, title in descriptions:
        sheet = book.create_sheet(sheet_name)
        sheet.append(["Архів", archive])
        sheet.append(["Фонд", fond])
        sheet.append(["Опис", inventory])
        sheet.append(["№ з/п", "Заголовок справи", "Крайні дати", "Кількість аркушів", "Примітки"])
        sheet.append([None, "1850 рік"])
        sheet.append([1, title, "1850", 2])
    book.save(path)


class MultiWorkbookTests(unittest.TestCase):
    def test_chronology_accepts_other_centuries(self):
        self.assertEqual(f.parse_years("1640–1642", 5, "1", [], "опис"), (1640, 1642))
        self.assertEqual(f.parse_years("2020–2022", 6, "2", [], "опис"), (2020, 2022))
        self.assertEqual(f.detect_status("", "1640 рік"), "year_heading")

    def test_missing_metadata_identifies_sheet_and_cell(self):
        with tempfile.TemporaryDirectory() as temp:
            path = Path(temp) / "broken.xlsx"
            workbook(path, [("опис", "ДАМО", 230, 1, "Справа про школу")])
            book = openpyxl.load_workbook(path)
            book.active["B2"] = None
            book.save(path)
            book.close()
            previous = f.INPUT_FILE
            try:
                f.INPUT_FILE = path
                with self.assertRaisesRegex(ValueError, "broken.xlsx, опис, B2"):
                    f.configure_workbook_sources(openpyxl)
            finally:
                f.INPUT_FILE = previous

    def test_metadata_parsed_and_exported_with_original_rows(self):
        with tempfile.TemporaryDirectory() as temp:
            path = Path(temp) / "sources.xlsx"
            workbook(path, [
                ("Невідомий 1", "ДАМО", 230, 1, "Справа про лікарню"),
                ("Інший опис", "ЦДІАК України", 356, 1, "Дело о больнице"),
            ])
            old = f.INPUT_FILE, f.SOURCE_CONFIG, f.SHEET_NAMES, f.HEADER_ROWS_BY_SHEET
            try:
                f.INPUT_FILE = path
                f.configure_workbook_sources(openpyxl)
                records, issues, _ = f.read_records(openpyxl)
                cases = [r for r in records if r.status == "case"]
                self.assertEqual([(r.excel_row, r.source_archive, r.source_fond, r.source_inventory)
                                  for r in cases], [(6, "ДАМО", "230", "1"), (6, "ЦДІАК України", "356", "1")])
                self.assertEqual([r.language for r in cases], ["uk", "ru"])
                self.assertNotEqual(cases[0].source_id, cases[1].source_id)
                self.assertFalse(any(i.field == "Структура рядка" for i in issues))
                source = openpyxl.load_workbook(path)
                output = Path(temp) / "filtered.xlsx"
                try:
                    f.create_filtered_workbook(openpyxl, source, cases[:1], output, 5)
                finally:
                    source.close()
                result = openpyxl.load_workbook(output)
                try:
                    first = result["Невідомий 1"]
                    self.assertEqual([first[f"A{i}"].value for i in (1, 2, 3, 4, 5)],
                                     ["Архів", "Фонд", "Опис", "№ з/п", 1])
                    self.assertEqual(first["B5"].value, "Справа про лікарню")
                    self.assertEqual(result["Інший опис"]["B3"].value, 1)
                    self.assertEqual(result["Інший опис"].max_row, 4)
                finally:
                    result.close()
            finally:
                f.INPUT_FILE, f.SOURCE_CONFIG, f.SHEET_NAMES, f.HEADER_ROWS_BY_SHEET = old

    def test_batch_isolates_books_and_keeps_processing_after_bad_book(self):
        with tempfile.TemporaryDirectory() as temp:
            base = Path(temp)
            directory = base / "input"
            directory.mkdir()
            workbook(directory / "A.xlsx", [
                ("one", "ДАМО", 230, 1, "Справа про лікарню"),
                ("two", "ЦДІАК України", 356, 1, "Дело о больнице"),
            ])
            (directory / "B.xlsx").write_bytes(b"broken")
            workbook(directory / "C.xlsx", [
                ("third", "ДАМО", 229, 1, "Справа про школу"),
            ])
            (directory / "~$ignored.xlsx").write_bytes(b"lock")
            old = f.INPUT_DIR, f.INPUT_FILE, f.OUTPUTS_DIR, f.ANALYSIS_CONFIG
            try:
                f.INPUT_DIR = directory
                f.OUTPUTS_DIR = base / "outputs"
                f.ANALYSIS_CONFIG = {"exports": {"enabled": False}}
                self.assertEqual([p.name for p in f.input_workbooks()], ["A.xlsx", "B.xlsx", "C.xlsx"])
                with patch.object(f, "configure_analysis", lambda _: None), \
                     patch.object(f, "load_dependencies", lambda: (openpyxl, yaml, None)), \
                     contextlib.redirect_stdout(io.StringIO()):
                    with self.assertRaisesRegex(RuntimeError, "Частина книг"):
                        f.main()
                run = next(f.OUTPUTS_DIR.iterdir())
                with (run / "run_summary.csv").open(encoding="utf-8-sig", newline="") as stream:
                    summary = list(csv.DictReader(stream, delimiter=";"))
                self.assertEqual([(row["file"], row["status"], row["cases"]) for row in summary],
                                 [("A.xlsx", "complete", "2"), ("B.xlsx", "failed", ""),
                                  ("C.xlsx", "complete", "1")])
                self.assertTrue((run / "B" / "failure.txt").exists())
                a_manifest = json.loads((run / "A" / "run_manifest.json").read_text())
                c_manifest = json.loads((run / "C" / "run_manifest.json").read_text())
                self.assertEqual(len(a_manifest["sources"]), 2)
                self.assertEqual(len(c_manifest["sources"]), 1)
                self.assertEqual(next(row["titles"] for row in a_manifest["source_summary"]
                                      if row["source"] == "combined"), 2)
                self.assertEqual(next(row["titles"] for row in c_manifest["source_summary"]
                                      if row["source"] == "combined"), 1)
            finally:
                f.INPUT_DIR, f.INPUT_FILE, f.OUTPUTS_DIR, f.ANALYSIS_CONFIG = old


if __name__ == "__main__":
    unittest.main()
