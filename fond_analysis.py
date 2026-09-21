from __future__ import annotations

import csv
import hashlib
import json
import platform
from copy import copy
import re
import statistics
import sys
import time
import traceback
from collections import Counter, defaultdict
from datetime import datetime
from pathlib import Path
from typing import Any, Iterable


from models import Issue, Category, Record, CATEGORY_RULE_FIELDS
from text_matching import (
    MATCH_LANGUAGE,
    WORD_RE,
    YEAR_HEADING_RE,
    YEAR_RE,
    LONG_NUMBER_RE,
    WITHDRAWN_RE,
    STEM_ENDINGS,
    clean_cell,
    normalize_text,
    tokenize,
    UK_IRREGULAR,
    RU_IRREGULAR,
    light_stem_word,
    ukrainian_stem_word,
    stem_tokens,
    cached_item_tokens,
    item_tokens,
    find_phrase_positions,
    contains_item,
    item_matcher,
    mask_item,
    normalize_case_id,
    RU_ENDINGS,
    russian_stem_word,
    detect_title_language,
)

BASE_DIR = Path(__file__).resolve().parent
SCRIPT_VERSION = "0.8"
CONFIG_DIR = BASE_DIR / "config"
INPUT_FILE = BASE_DIR / "input.xlsx"
DICTIONARIES_DIR = BASE_DIR / "dictionaries"
OUTPUTS_DIR = BASE_DIR / "outputs"
RUN_DIR = OUTPUTS_DIR
REPORTS_DIR = RUN_DIR / "reports"
FIGURES_DIR = RUN_DIR / "figures"
WORK_DIR = RUN_DIR / "tables"
THEMATIC_DIR = RUN_DIR / "thematic_exports"

SHEET_NAMES = ("Опис 1", "Опис 2", "Опис 3", "Опис 4", "ЦДІАК")
SOURCE_CONFIG: dict[str, dict[str, Any]] = {}
ANALYSIS_CONFIG: dict[str, Any] = {}
LANGUAGE_CATEGORIES: dict[str, list] = {}
LANGUAGE_AMBIGUITIES: dict[str, list] = {}
MIN_ALLOWED_YEAR = 1790
MAX_ALLOWED_YEAR = 1911


def load_dependencies():
    missing: list[str] = []
    try:
        import openpyxl
    except ImportError:
        openpyxl = None
        missing.append("openpyxl")
    try:
        import yaml
    except ImportError:
        yaml = None
        missing.append("pyyaml")

    if missing:
        print("Не встановлено обов'язкові бібліотеки:")
        print("  " + ", ".join(missing))
        print("\nВстановіть їх у терміналі PyCharm:")
        print(f'  "{sys.executable}" -m pip install ' + " ".join(missing))
        raise SystemExit(1)

    try:
        import matplotlib
        matplotlib.use("Agg")
        import matplotlib.pyplot as plt
    except ImportError:
        plt = None
    return openpyxl, yaml, plt


def config_path(name: str, legacy: Path | None = None) -> Path:
    """Повертає канонічний конфіг, залишаючи читання старої структури."""
    preferred = CONFIG_DIR / name
    if preferred.exists():
        return preferred
    if legacy is not None and legacy.exists():
        return legacy
    return preferred


def configure_analysis(yaml_module) -> None:
    global ANALYSIS_CONFIG, INPUT_FILE, OUTPUTS_DIR
    global MIN_ALLOWED_YEAR, MAX_ALLOWED_YEAR
    path = config_path("analysis.yaml")
    config = safe_load_yaml(path, yaml_module)
    ANALYSIS_CONFIG = config

    input_name = str(config.get("input_file", "input.xlsx"))
    output_name = str(config.get("output_root", "outputs"))
    input_path = Path(input_name)
    output_path = Path(output_name)
    INPUT_FILE = input_path if input_path.is_absolute() else BASE_DIR / input_path
    OUTPUTS_DIR = output_path if output_path.is_absolute() else BASE_DIR / output_path

    chronology = config.get("chronology", {})
    MIN_ALLOWED_YEAR = int(chronology.get("minimum_year", 1790))
    MAX_ALLOWED_YEAR = int(chronology.get("maximum_year", 1911))
    if MIN_ALLOWED_YEAR > MAX_ALLOWED_YEAR:
        raise ValueError("config/analysis.yaml: minimum_year більший за maximum_year")


def configure_sources(yaml_module) -> None:
    global SOURCE_CONFIG, SHEET_NAMES
    source_path = config_path("sources.yaml", BASE_DIR / "sources.yaml")
    config = safe_load_yaml(source_path, yaml_module)
    SOURCE_CONFIG = config.get("sheets", {})
    if not SOURCE_CONFIG:
        raise ValueError("sources.yaml: немає налаштованих аркушів")
    for name, item in SOURCE_CONFIG.items():
        if item.get("language") not in {"uk", "ru"}:
            raise ValueError(f"Непідтримувана мова: {name}")
        if not re.fullmatch(r"[a-z0-9_-]+", item.get("source_id", "")):
            raise ValueError(f"Небезпечний або порожній source_id: {name}")
    SHEET_NAMES = tuple(SOURCE_CONFIG)


def safe_load_yaml(path: Path, yaml_module) -> dict[str, Any]:
    with path.open("r", encoding="utf-8") as stream:
        data = yaml_module.safe_load(stream)
    if not isinstance(data, dict):
        raise ValueError(f"Корінь YAML-файла має бути словником: {path}")
    return data


def ensure_directories() -> None:
    REPORTS_DIR.mkdir(parents=True, exist_ok=True)
    FIGURES_DIR.mkdir(parents=True, exist_ok=True)
    WORK_DIR.mkdir(parents=True, exist_ok=True)
    THEMATIC_DIR.mkdir(parents=True, exist_ok=True)


def load_dictionaries(yaml_module, language="uk"):
    index_path = config_path(
        "categories.yaml", DICTIONARIES_DIR / "categories.yaml"
    )
    if not index_path.exists():
        raise FileNotFoundError(
            f"Не знайдено {index_path}. Папка dictionaries має лежати поруч зі скриптом."
        )

    index = safe_load_yaml(index_path, yaml_module)
    categories: list[Category] = []
    configured_categories = index.get("categories", [])
    if not isinstance(configured_categories, list):
        raise ValueError(f"{index_path}: categories має бути списком")
    macroblock_config = index.get("macroblocks", {})
    category_ids: set[str] = set()
    export_names: set[str] = set()
    for item in configured_categories:
        if not isinstance(item, dict):
            raise ValueError(f"{index_path}: кожна категорія має бути словником")
        if not isinstance(item.get("enabled", True), bool):
            raise ValueError(f"{index_path}: enabled має бути true або false")
        if not item.get("enabled", True):
            continue
        category_id = str(item.get("id", ""))
        if not re.fullmatch(r"[a-z0-9_]+", category_id):
            raise ValueError(f"Некоректний id категорії: {category_id}")
        if category_id in category_ids:
            raise ValueError(f"Категорія повторюється: {category_id}")
        category_ids.add(category_id)
        if item.get("macroblock") not in macroblock_config:
            raise ValueError(
                f"Невідомий макроблок для {category_id}: {item.get('macroblock')}"
            )
        dictionaries = item.get("dictionaries", {})
        configured_path = dictionaries.get(language) if isinstance(dictionaries, dict) else None
        if configured_path:
            category_path = BASE_DIR / str(configured_path)
        else:
            # Сумісність із categories.yaml версій 0.3–0.6.
            category_path = (
                DICTIONARIES_DIR
                / ("ru" if language == "ru" else "")
                / item["file"]
            )
        data = safe_load_yaml(category_path, yaml_module)
        if data.get("id") != item["id"]:
            raise ValueError(f"Ідентифікатор категорії не відповідає індексу: {category_path}")
        for key in CATEGORY_RULE_FIELDS:
            if key == "context_rules":
                continue
            if not isinstance(data.get(key, []), list) or not all(isinstance(v, str) for v in data.get(key, [])):
                raise ValueError(f"{category_path}: {key} має бути списком рядків")
        if not isinstance(data.get("context_rules", []), list) or not all(
            isinstance(rule, dict) for rule in data.get("context_rules", [])
        ):
            raise ValueError(f"{category_path}: context_rules має бути списком правил")
        for rule in data.get("context_rules", []):
            for key in ("all", "any"):
                if key in rule and (
                    not isinstance(rule[key], list)
                    or not all(isinstance(value, str) for value in rule[key])
                ):
                    raise ValueError(
                        f"{category_path}: {key} у context_rules має бути списком рядків"
                    )
        export_filename = str(
            item.get("export", f"{item['id']}.xlsx")
        ).strip()
        if not re.fullmatch(r'[^<>:"/\\|?*]+\.xlsx', export_filename, re.I):
            raise ValueError(
                f"Некоректна назва тематичного файла: {export_filename}"
            )
        export_key = export_filename.casefold()
        if export_key in export_names:
            raise ValueError(
                f"Назва тематичного файла повторюється: {export_filename}"
            )
        export_names.add(export_key)
        categories.append(
            Category(
                id=data["id"],
                label=item["label"],
                macroblock=data["macroblock"],
                minimum_score=int(
                    data.get(
                        "minimum_score",
                        index.get("classification", {}).get("minimum_score", 3),
                    )
                ),
                **{key: list(data.get(key, [])) for key in CATEGORY_RULE_FIELDS},
                export_filename=export_filename,
            )
        )
    if not categories:
        raise ValueError("У categories.yaml немає тематичних категорій.")

    ambiguity_path = DICTIONARIES_DIR / ("ru" if language == "ru" else "") / "auxiliary" / "ambiguities.yaml"
    ambiguities = (
        safe_load_yaml(ambiguity_path, yaml_module).get("rules", [])
        if ambiguity_path.exists()
        else []
    )

    macroblocks = {
        key: value.get("label", key)
        for key, value in index.get("macroblocks", {}).items()
    }
    dictionary_version = str(index.get("dictionary_version", "невідома"))
    return categories, ambiguities, macroblocks, dictionary_version


def detect_status(
    case_id: str,
    title: str,
    dates: str = "",
    pages: str = "",
    notes: str = "",
) -> str:
    title_normalized = normalize_text(title).strip()
    if YEAR_HEADING_RE.match(title_normalized) and not case_id:
        return "year_heading"
    # У частині описів службову позначку надруковано з міжлітерними
    # пробілами: «В И Б У Л А», «В И Б У Л И», «В И Б У Л О».
    if WITHDRAWN_RE.match(title_normalized):
        return "withdrawn"
    if re.match(r"^(?:архівний опис|архивная опись|недействующая опись|недіючий опис)(?:\s|$|[.,])", title_normalized):
        return "inventory"
    if case_id and title:
        return "case"
    # Загальні заголовки груп справ можуть не мати номера.
    # Якщо інші облікові поля такого рядка порожні, це не помилка.
    if title and not case_id and not any((dates, pages, notes)):
        return "group_heading"
    if case_id or title or dates or pages or notes:
        return "unknown"
    return "empty"


def parse_years(
    dates_raw: str,
    row: int,
    case_id: str,
    issues: list[Issue],
    sheet_name: str,
) -> tuple[int | None, int | None]:
    if not dates_raw:
        issues.append(
            Issue(
                "WARNING", row, case_id, "Крайні дати", "",
                "Для справи не зазначено крайні дати.",
                sheet_name,
            )
        )
        return None, None

    long_numbers = LONG_NUMBER_RE.findall(dates_raw)
    if long_numbers:
        issues.append(
            Issue(
                "ERROR", row, case_id, "Крайні дати", dates_raw,
                "Виявлено число з п'ятьма або більше цифрами: "
                + ", ".join(long_numbers),
                sheet_name,
            )
        )

    years = [int(value) for value in YEAR_RE.findall(dates_raw)]
    if not years:
        issues.append(
            Issue(
                "WARNING", row, case_id, "Крайні дати", dates_raw,
                "Не вдалося розпізнати чотиризначний рік.",
                sheet_name,
            )
        )
        return None, None

    for year in sorted(set(years)):
        if year < MIN_ALLOWED_YEAR or year > MAX_ALLOWED_YEAR:
            issues.append(
                Issue(
                    "WARNING", row, case_id, "Крайні дати", dates_raw,
                    f"Рік {year} виходить за контрольні межі "
                    f"{MIN_ALLOWED_YEAR}–{MAX_ALLOWED_YEAR}.",
                    sheet_name,
                )
            )

    start_year = years[0]
    end_year = years[-1]
    if end_year < start_year:
        issues.append(
            Issue(
                "ERROR", row, case_id, "Крайні дати", dates_raw,
                f"Кінцевий рік {end_year} менший за початковий {start_year}.",
                sheet_name,
            )
        )
    return start_year, end_year


def parse_pages(
    pages_raw: str,
    row: int,
    case_id: str,
    issues: list[Issue],
    sheet_name: str,
) -> int | None:
    if not pages_raw:
        issues.append(
            Issue(
                "WARNING", row, case_id, "Кількість аркушів", "",
                "Не зазначено кількість аркушів.",
                sheet_name,
            )
        )
        return None

    compact = re.sub(r"\s+", "", pages_raw)
    if re.fullmatch(r"\d+", compact):
        pages = int(compact)
        if pages <= 0:
            issues.append(
                Issue(
                    "WARNING", row, case_id, "Кількість аркушів", pages_raw,
                    "Кількість аркушів має бути більшою за нуль.",
                    sheet_name,
                )
            )
            return None
        return pages

    numbers = re.findall(r"\d+", compact)
    if len(numbers) == 1:
        issues.append(
            Issue(
                "WARNING", row, case_id, "Кількість аркушів", pages_raw,
                "Нетипове оформлення; використано знайдене числове значення.",
                sheet_name,
            )
        )
        return int(numbers[0])

    issues.append(
        Issue(
            "WARNING", row, case_id, "Кількість аркушів", pages_raw,
            "Не вдалося однозначно визначити кількість аркушів.",
            sheet_name,
        )
    )
    return None


def read_records(
    openpyxl_module,
) -> tuple[list[Record], list[Issue], dict[str, list[str]]]:
    if not INPUT_FILE.exists():
        raise FileNotFoundError(
            f"Не знайдено {INPUT_FILE}. Покладіть input.xlsx поруч зі скриптом."
        )

    workbook = openpyxl_module.load_workbook(
        INPUT_FILE,
        read_only=True,
        data_only=True,
    )
    available_sheets = [name for name in SHEET_NAMES if name in workbook.sheetnames]
    if not available_sheets:
        available = ", ".join(workbook.sheetnames)
        raise KeyError(
            "У книзі немає жодного з очікуваних аркушів "
            f"{', '.join(SHEET_NAMES)}. Доступні аркуші: {available}"
        )

    issues: list[Issue] = []
    records: list[Record] = []
    expected_headers = [
        "№", "заголовок", "крайні дати", "кількість аркушів", "примітки"
    ]
    headers_by_sheet: dict[str, list[str]] = {}
    for ignored in set(workbook.sheetnames) - set(SHEET_NAMES):
        issues.append(Issue("WARNING", None, "", "Аркуш", ignored,
                            "Аркуш не налаштований у sources.yaml і не включений до аналізу.", ignored))

    for missing_sheet in (name for name in SHEET_NAMES if name not in workbook.sheetnames):
        issues.append(
            Issue(
                "WARNING", None, "", "Аркуш", missing_sheet,
                "Очікуваний аркуш відсутній у книзі.",
                missing_sheet,
            )
        )

    for sheet_name in available_sheets:
        worksheet = workbook[sheet_name]
        current_section_year: int | None = None
        actual_headers = [
            normalize_text(clean_cell(cell.value))
            for cell in next(
                worksheet.iter_rows(min_row=1, max_row=1, max_col=5)
            )
        ]
        headers_by_sheet[sheet_name] = actual_headers

        config = SOURCE_CONFIG.get(sheet_name, {})
        current_section_label = ""
        for column, expected in enumerate(expected_headers, start=1):
            if column in config.get("skip_header_checks", []):
                continue
            if expected not in actual_headers[column - 1]:
                issues.append(
                    Issue(
                        "WARNING", 1, "", f"Колонка {column}",
                        actual_headers[column - 1],
                        f"Заголовок колонки не містить очікуваний текст «{expected}».",
                        sheet_name,
                    )
                )

        for excel_row, values in enumerate(
            worksheet.iter_rows(min_row=2, max_col=5, values_only=True),
            start=2,
        ):
            case_id_raw, title_raw, dates_raw, pages_raw, notes_raw = [
                clean_cell(value) for value in values
            ]
            status = detect_status(
                case_id_raw, title_raw, dates_raw, pages_raw, notes_raw
            )
            if status == "empty":
                continue
            if status == "year_heading":
                heading_years = [
                    int(year) for year in YEAR_RE.findall(normalize_text(title_raw))
                ]
                current_section_year = (
                    heading_years[0] if len(heading_years) == 1 else None
                )
                current_section_label = title_raw
            elif status == "group_heading":
                current_section_label = title_raw
                current_section_year = None

            record = Record(
                excel_row=excel_row,
                sheet_name=sheet_name,
                case_id_raw=case_id_raw,
                case_id_normalized=normalize_case_id(case_id_raw),
                title_raw=title_raw,
                dates_raw=dates_raw,
                pages_raw=pages_raw,
                notes_raw=notes_raw,
                status=status,
                section_year=current_section_year,
                section_label=current_section_label,
                source_archive=config.get("archive", ""),
                source_fond=str(config.get("fond", "")),
                source_inventory=str(config.get("inventory", "")),
                source_id=config.get("source_id", sheet_name),
                language=config.get("language", "uk"),
                detected_language=detect_title_language(title_raw),
            )

            if status == "case":
                record.start_year, record.end_year = parse_years(
                    dates_raw, excel_row, case_id_raw, issues, sheet_name
                )
                record.pages = parse_pages(
                    pages_raw, excel_row, case_id_raw, issues, sheet_name
                ) if pages_raw or config.get("pages_expected", True) else None
            elif status == "unknown":
                joined = " | ".join(
                    value for value in (
                        case_id_raw, title_raw, dates_raw, pages_raw, notes_raw
                    ) if value
                )
                issues.append(
                    Issue(
                        "ERROR", excel_row, case_id_raw, "Структура рядка", joined,
                        "Рядок не вдалося віднести до відомого типу.",
                        sheet_name,
                    )
                )
            records.append(record)

    workbook.close()

    occurrences: dict[tuple[str, str], list[Record]] = defaultdict(list)
    for record in records:
        if record.case_id_normalized and record.status in {"case", "withdrawn"}:
            key = (record.sheet_name, record.case_id_normalized)
            occurrences[key].append(record)

    for (sheet_name, normalized_id), duplicated in occurrences.items():
        if len(duplicated) <= 1:
            continue
        rows = ", ".join(str(record.excel_row) for record in duplicated)
        titles = " || ".join(
            f"рядок {record.excel_row}: {record.title_raw}"
            for record in duplicated
        )
        issues.append(
            Issue(
                "WARNING", duplicated[0].excel_row, duplicated[0].case_id_raw,
                "Номер справи", normalized_id,
                f"Номер повторюється у рядках {rows}. {titles}",
                sheet_name,
            )
        )

    return records, issues, headers_by_sheet


def global_masks_for_category(
    tokens: list[str],
    category_id: str,
    ambiguities: list[dict[str, Any]],
) -> list[str]:
    result = list(tokens)
    for rule in ambiguities:
        term = str(rule.get("term", ""))
        if category_id in rule.get("do_not_assign", []):
            result = mask_item(result, term)
        key = f"exclude_for_{category_id}"
        for phrase in rule.get(key, []):
            result = mask_item(result, str(phrase))
    return result


def check_context_rule(tokens: list[str], rule: dict[str, Any], matches=None) -> bool:
    matches = matches if matches is not None else item_matcher(tokens)
    all_items = [str(value) for value in rule.get("all", [])]
    any_items = [str(value) for value in rule.get("any", [])]
    if all_items and not all(matches(item) for item in all_items):
        return False
    if any_items and not any(matches(item) for item in any_items):
        return False
    return bool(all_items or any_items)


def classify_title(
    title: str,
    categories: list[Category],
    ambiguities: list[dict[str, Any]],
    macroblock_order: list[str],
    language: str = "uk",
):
    token = MATCH_LANGUAGE.set(language)
    try:
        return _classify_title(
            title, LANGUAGE_CATEGORIES.get(language, categories),
            LANGUAGE_AMBIGUITIES.get(language, ambiguities), macroblock_order
        )
    finally:
        MATCH_LANGUAGE.reset(token)


def _classify_title(title, categories, ambiguities, macroblock_order):
    original_tokens = stem_tokens(title)
    matched_ids: list[str] = []
    matched_labels: list[str] = []
    scores: dict[str, int] = {}
    evidence: dict[str, list[str]] = {}
    context_ids: list[str] = []
    context_labels: list[str] = []
    context_evidence: dict[str, list[str]] = {}
    review_flags: set[str] = set()

    for category in categories:
        tokens = global_masks_for_category(original_tokens, category.id, ambiguities)
        for phrase in category.exclude_phrases:
            tokens = mask_item(tokens, phrase)

        matches = item_matcher(tokens)
        category_context_evidence: list[str] = []
        for items, label in (
            (category.context_only_phrases, "фраза"),
            (category.context_only_terms, "термін"),
            (category.contextual_terms, "слабкий контекст"),
        ):
            category_context_evidence.extend(
                f"{label}: {item}" for item in items if matches(item)
            )
        if category_context_evidence:
            context_ids.append(category.id)
            context_labels.append(category.label)
            context_evidence[category.id] = category_context_evidence

        score = 0
        category_evidence: list[str] = []
        for items, weight, label in (
            (category.strong_phrases, 4, "фраза"),
            (category.strong_terms, 3, "термін"),
            (category.contextual_terms, 1, "контекст"),
        ):
            for item in items:
                if matches(item):
                    score += weight
                    category_evidence.append(f"{label}: {item}")
        for rule in category.context_rules:
            if check_context_rule(tokens, rule, matches):
                score += int(rule.get("score", 3))
                matched_rule_items = [
                    str(value) for value in rule.get("all", [])
                    if matches(str(value))
                ] + [
                    str(value) for value in rule.get("any", [])
                    if matches(str(value))
                ]
                readable = " + ".join(matched_rule_items)
                category_evidence.append(f"правило: {readable}")

        for ambiguity in ambiguities:
            term = str(ambiguity.get("term", ""))
            category_context = ambiguity.get("conditional", {}).get(category.id, [])
            if (
                term
                and category_context
                and matches(term)
                and any(
                    matches(str(context))
                    for context in category_context
                )
            ):
                score += 3
                category_evidence.append(
                    f"контекст неоднозначного терміна: {term}"
                )

        has_subject_evidence = any(not e.startswith("контекст:") for e in category_evidence)
        if score >= category.minimum_score and (MATCH_LANGUAGE.get() != "ru" or has_subject_evidence):
            matched_ids.append(category.id)
            matched_labels.append(category.label)
            scores[category.id] = score
            evidence[category.id] = category_evidence
        elif score >= max(1, category.minimum_score - 1):
            for review_term in category.review_terms:
                if matches(review_term):
                    review_flags.add(f"{category.id}: {review_term}")

    ordered = sorted(
        zip(matched_ids, matched_labels),
        key=lambda pair: (-scores[pair[0]], pair[1]),
    )
    matched_ids = [pair[0] for pair in ordered]
    matched_labels = [pair[1] for pair in ordered]
    category_by_id = {category.id: category for category in categories}
    found_macroblocks = {
        category_by_id[category_id].macroblock for category_id in matched_ids
    }
    macroblocks = [
        macroblock for macroblock in macroblock_order
        if macroblock in found_macroblocks
    ]
    return (
        matched_ids, matched_labels, macroblocks,
        scores, evidence, context_ids, context_labels,
        context_evidence, sorted(review_flags)
    )


def classify_records(
    records: list[Record],
    categories: list[Category],
    ambiguities: list[dict[str, Any]],
    macroblock_order: list[str],
) -> None:
    active_records = [record for record in records if record.status == "case"]
    total = len(active_records)
    if total == 0:
        print("   Немає справ для тематичної класифікації.")
        return

    started_at = time.perf_counter()
    update_every = max(1, total // 100)
    print_classification_progress(0, total, started_at)

    for processed, record in enumerate(active_records, start=1):
        (
            record.categories,
            record.category_labels,
            record.macroblocks,
            record.scores,
            record.evidence,
            record.context_categories,
            record.context_category_labels,
            record.context_evidence,
            record.review_flags,
        ) = classify_title(
            record.title_raw, categories, ambiguities, macroblock_order,
            language=record.language
        )
        if record.detected_language not in {"undetermined", record.language}:
            record.review_flags.append("language_check: " + record.detected_language)
        if processed % update_every == 0 or processed == total:
            print_classification_progress(
                processed, total, started_at, finished=processed == total
            )


def format_duration(seconds: float) -> str:
    total_seconds = max(0, int(round(seconds)))
    hours, remainder = divmod(total_seconds, 3600)
    minutes, seconds = divmod(remainder, 60)
    if hours:
        return f"{hours:02d}:{minutes:02d}:{seconds:02d}"
    return f"{minutes:02d}:{seconds:02d}"


def format_count(value: int) -> str:
    return f"{value:,}".replace(",", " ")


def print_classification_progress(
    processed: int,
    total: int,
    started_at: float,
    finished: bool = False,
) -> None:
    elapsed = time.perf_counter() - started_at
    fraction = processed / total if total else 1.0
    bar_width = 28
    filled = min(bar_width, int(bar_width * fraction))
    bar = "█" * filled + "·" * (bar_width - filled)

    details = (
        f"   [{bar}] {fraction * 100:6.2f}% | "
        f"{format_count(processed)}/{format_count(total)} | "
        f"час {format_duration(elapsed)}"
    )
    if processed:
        remaining = elapsed / processed * (total - processed)
        if finished:
            details += " | завершено"
        else:
            details += f" | залишилося ~{format_duration(remaining)}"

    print("\r" + details.ljust(115), end="\n" if finished else "", flush=True)


def write_csv(path: Path, fieldnames: list[str],
              rows: Iterable[dict[str, Any]]) -> int:
    path.parent.mkdir(parents=True, exist_ok=True)
    written = 0
    with path.open("w", encoding="utf-8-sig", newline="") as stream:
        writer = csv.DictWriter(
            stream, fieldnames=fieldnames, delimiter=";", extrasaction="ignore"
        )
        writer.writeheader()
        for row in rows:
            writer.writerow(row)
            written += 1
    return written


def classification_type(record: Record) -> str:
    if len(record.categories) > 1:
        return "subject_multilabel"
    if record.categories:
        return "subject_single"
    if record.context_categories:
        return "context_only"
    return "unclassified"


def classification_confidence(record: Record) -> str:
    """Описова сила правила, а не статистична ймовірність."""
    if record.categories:
        return "strong" if max(record.scores.values(), default=0) >= 4 else "threshold"
    if record.context_categories:
        return "context_only"
    return "none"


CLASSIFICATION_INDEX_FIELDS = [
    "record_uid", "source", "archive", "fond", "inventory", "sheet",
    "row_number", "case_id", "status", "title_original", "dates_original",
    "pages_original", "notes_original", "language", "detected_language",
    "categories", "category_labels", "classification_type",
    "classification_confidence", "scores", "evidence", "context_categories",
    "context_labels", "context_evidence", "section_label", "start_year",
    "end_year", "chronology_basis",
]


def classification_index_row(record: Record) -> dict[str, Any]:
    base = record_to_row(record)
    return {
        "record_uid": base["record_uid"],
        "source": record.source_id,
        "archive": record.source_archive,
        "fond": record.source_fond,
        "inventory": record.source_inventory,
        "sheet": record.sheet_name,
        "row_number": record.excel_row,
        "case_id": record.case_id_raw,
        "status": record.status,
        "title_original": record.title_raw,
        "dates_original": record.dates_raw,
        "pages_original": record.pages_raw,
        "notes_original": record.notes_raw,
        "language": record.language,
        "detected_language": record.detected_language,
        "categories": " | ".join(record.categories),
        "category_labels": " | ".join(record.category_labels),
        "classification_type": classification_type(record),
        "classification_confidence": classification_confidence(record),
        "scores": base["scores"],
        "evidence": base["evidence"],
        "context_categories": " | ".join(record.context_categories),
        "context_labels": " | ".join(record.context_category_labels),
        "context_evidence": base["context_evidence"],
        "section_label": record.section_label,
        "start_year": record.start_year or "",
        "end_year": record.end_year or "",
        "chronology_basis": chronology_basis(record),
    }


def write_combined_indexes(
    records: list[Record], categories: list[Category], tables_dir: Path
) -> None:
    cases = [record for record in records if record.status == "case"]
    unclassified = [record for record in cases if not record.categories]
    service_records = [record for record in records if record.status != "case"]

    written = write_csv(
        tables_dir / "classification_index.csv",
        CLASSIFICATION_INDEX_FIELDS,
        (classification_index_row(record) for record in cases),
    )
    if written != len(cases):
        raise RuntimeError(
            "classification_index.csv: записано "
            f"{written} рядків замість {len(cases)}"
        )
    written = write_csv(
        tables_dir / "unclassified_cases.csv",
        CLASSIFICATION_INDEX_FIELDS,
        (classification_index_row(record) for record in unclassified),
    )
    if written != len(unclassified):
        raise RuntimeError(
            "unclassified_cases.csv: записано "
            f"{written} рядків замість {len(unclassified)}"
        )

    service_fields = [
        "record_uid", "source", "archive", "fond", "inventory", "sheet",
        "row_number", "case_id", "title_original", "dates_original",
        "pages_original", "notes_original", "status", "section_label",
    ]
    write_csv(
        tables_dir / "service_records.csv",
        service_fields,
        (
            {
                "record_uid": f"{r.source_id}:{r.sheet_name}:{r.excel_row}",
                "source": r.source_id,
                "archive": r.source_archive,
                "fond": r.source_fond,
                "inventory": r.source_inventory,
                "sheet": r.sheet_name,
                "row_number": r.excel_row,
                "case_id": r.case_id_raw,
                "title_original": r.title_raw,
                "dates_original": r.dates_raw,
                "pages_original": r.pages_raw,
                "notes_original": r.notes_raw,
                "status": r.status,
                "section_label": r.section_label,
            }
            for r in service_records
        ),
    )

    counts = Counter(
        category_id for record in cases for category_id in record.categories
    )
    write_csv(
        tables_dir / "thematic_summary.csv",
        [
            "category_id", "category_label", "macroblock", "export_file",
            "cases", "percent_of_analyzed_titles",
        ],
        (
            {
                "category_id": category.id,
                "category_label": category.label,
                "macroblock": category.macroblock,
                "export_file": category.export_filename,
                "cases": counts[category.id],
                "percent_of_analyzed_titles": (
                    f"{counts[category.id] / len(cases) * 100:.2f}"
                    if cases else "0.00"
                ),
            }
            for category in categories
        ),
    )


def copy_cell(source_cell, target_cell) -> None:
    target_cell.value = source_cell.value
    # Внутрішній _style містить індекси, дійсні лише для книги-джерела.
    # У новій книзі копіюємо складові стилю окремо.
    target_cell.font = copy(source_cell.font)
    target_cell.fill = copy(source_cell.fill)
    target_cell.border = copy(source_cell.border)
    target_cell.alignment = copy(source_cell.alignment)
    target_cell.number_format = source_cell.number_format
    target_cell.protection = copy(source_cell.protection)
    target_cell.quotePrefix = source_cell.quotePrefix
    if source_cell.hyperlink:
        target_cell._hyperlink = copy(source_cell.hyperlink)
    if source_cell.comment:
        target_cell.comment = copy(source_cell.comment)


def copy_source_row(source_sheet, target_sheet, source_row: int,
                    target_row: int, max_columns: int) -> None:
    for column in range(1, max_columns + 1):
        copy_cell(
            source_sheet.cell(source_row, column),
            target_sheet.cell(target_row, column),
        )
    source_dimension = source_sheet.row_dimensions[source_row]
    target_dimension = target_sheet.row_dimensions[target_row]
    if source_dimension.height is not None:
        target_dimension.height = source_dimension.height
    target_dimension.hidden = source_dimension.hidden
    target_dimension.outlineLevel = source_dimension.outlineLevel
    target_dimension.collapsed = source_dimension.collapsed
    target_dimension.thickTop = source_dimension.thickTop
    target_dimension.thickBot = source_dimension.thickBot


def create_filtered_workbook(
    openpyxl_module,
    source_workbook,
    records: list[Record],
    output_path: Path,
    max_columns: int,
) -> None:
    workbook = openpyxl_module.Workbook()
    workbook.remove(workbook.active)
    by_sheet: dict[str, list[Record]] = defaultdict(list)
    for record in records:
        by_sheet[record.sheet_name].append(record)

    for sheet_name in SHEET_NAMES:
        target_sheet = workbook.create_sheet(sheet_name)
        source_sheet = (
            source_workbook[sheet_name]
            if sheet_name in source_workbook.sheetnames else None
        )
        if source_sheet is not None:
            copy_source_row(source_sheet, target_sheet, 1, 1, max_columns)
            for column in range(1, max_columns + 1):
                letter = openpyxl_module.utils.get_column_letter(column)
                source_dimension = source_sheet.column_dimensions[letter]
                target_dimension = target_sheet.column_dimensions[letter]
                source_width = source_dimension.width
                if source_width is not None:
                    target_dimension.width = source_width
                target_dimension.hidden = source_dimension.hidden
                target_dimension.bestFit = source_dimension.bestFit
                target_dimension.outlineLevel = source_dimension.outlineLevel
                target_dimension.collapsed = source_dimension.collapsed
            target_sheet.sheet_view.showGridLines = source_sheet.sheet_view.showGridLines
            target_sheet.sheet_view.zoomScale = source_sheet.sheet_view.zoomScale
            target_sheet.sheet_view.zoomScaleNormal = (
                source_sheet.sheet_view.zoomScaleNormal
            )
            target_sheet.sheet_format.defaultRowHeight = (
                source_sheet.sheet_format.defaultRowHeight
            )
            target_sheet.sheet_format.defaultColWidth = (
                source_sheet.sheet_format.defaultColWidth
            )
            target_sheet.freeze_panes = source_sheet.freeze_panes
        else:
            fallback_headers = [
                "№ з/п", "Заголовок справи", "Крайні дати документів справи",
                "Кількість аркушів у справі", "Примітки",
            ]
            for column, value in enumerate(fallback_headers[:max_columns], start=1):
                target_sheet.cell(1, column, value)

        target_row = 2
        for record in sorted(by_sheet.get(sheet_name, []), key=lambda r: r.excel_row):
            if source_sheet is not None:
                copy_source_row(
                    source_sheet, target_sheet, record.excel_row,
                    target_row, max_columns,
                )
            else:
                values = [
                    record.case_id_raw, record.title_raw, record.dates_raw,
                    record.pages_raw, record.notes_raw,
                ]
                for column, value in enumerate(values[:max_columns], start=1):
                    target_sheet.cell(target_row, column, value)
            target_row += 1

        if (source_sheet is not None and source_sheet.auto_filter.ref
                and target_row > 2):
            last_column = openpyxl_module.utils.get_column_letter(max_columns)
            target_sheet.auto_filter.ref = f"A1:{last_column}{target_row - 1}"

    output_path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(output_path)
    workbook.close()


def export_thematic_selections(
    openpyxl_module, records: list[Record], categories: list[Category]
) -> list[dict[str, Any]]:
    exports = ANALYSIS_CONFIG.get("exports", {})
    if not exports.get("enabled", True):
        print("   Тематичний експорт вимкнено у config/analysis.yaml.")
        return []

    max_columns = int(exports.get("source_columns", 5))
    if not 1 <= max_columns <= 26:
        raise ValueError("exports.source_columns має бути в межах 1–26")
    create_xlsx = bool(exports.get("xlsx", True))
    create_csv = bool(exports.get("csv", True))
    if not create_xlsx and not create_csv:
        raise ValueError("У тематичному експорті потрібно ввімкнути XLSX або CSV")
    cases = [record for record in records if record.status == "case"]
    unclassified = [record for record in cases if not record.categories]
    service_records = [record for record in records if record.status != "case"]

    selections: list[tuple[str, str, list[Record]]] = [
        (category.id, category.export_filename,
         [r for r in cases if category.id in r.categories])
        for category in categories
    ]
    selections.extend(
        [
            (
                "unclassified",
                str(exports.get("unclassified_file", "некласифіковані.xlsx")),
                unclassified,
            ),
            (
                "service_records",
                str(exports.get("service_records_file", "службові_записи.xlsx")),
                service_records,
            ),
        ]
    )
    export_names: set[str] = set()
    for selection_id, filename, _ in selections:
        if not re.fullmatch(r'[^<>:"/\\|?*]+\.xlsx', filename, re.I):
            raise ValueError(
                f"Некоректна назва XLSX для {selection_id}: {filename}"
            )
        key = filename.casefold()
        if key in export_names:
            raise ValueError(f"Назва тематичного файла повторюється: {filename}")
        export_names.add(key)

    source_workbook = None
    if create_xlsx:
        source_workbook = openpyxl_module.load_workbook(
            INPUT_FILE, read_only=False, data_only=False
        )

    csv_dir = THEMATIC_DIR / "csv"
    manifest_rows: list[dict[str, Any]] = []
    total = len(selections)
    try:
        for index, (selection_id, filename, selected) in enumerate(selections, start=1):
            xlsx_path = THEMATIC_DIR / filename
            if create_xlsx:
                create_filtered_workbook(
                    openpyxl_module, source_workbook, selected,
                    xlsx_path, max_columns,
                )
            if create_csv:
                write_csv(
                    csv_dir / f"{Path(filename).stem}.csv",
                    CLASSIFICATION_INDEX_FIELDS,
                    (classification_index_row(record) for record in selected),
                )
            manifest_rows.append(
                {
                    "selection_id": selection_id,
                    "xlsx_file": filename if create_xlsx else "",
                    "csv_file": f"csv/{Path(filename).stem}.csv" if create_csv else "",
                    "records": len(selected),
                }
            )
            print(
                f"\r   Тематичні файли: {index}/{total} — {filename} "
                f"({format_count(len(selected))} записів)".ljust(115),
                end="\n" if index == total else "", flush=True,
            )
    finally:
        if source_workbook is not None:
            source_workbook.close()

    write_csv(
        THEMATIC_DIR / "export_manifest.csv",
        ["selection_id", "xlsx_file", "csv_file", "records"],
        manifest_rows,
    )
    return manifest_rows


def record_to_row(record: Record) -> dict[str, Any]:
    evidence_text = " | ".join(
        f"{category_id}: {', '.join(values)}"
        for category_id, values in record.evidence.items()
    )
    scores_text = " | ".join(
        f"{category_id}={score}" for category_id, score in record.scores.items()
    )
    context_evidence_text = " | ".join(
        f"{category_id}: {', '.join(values)}"
        for category_id, values in record.context_evidence.items()
    )
    return {
        "record_uid": f"{record.source_id}:{record.sheet_name}:{record.excel_row}",
        "source_archive": record.source_archive,
        "source_fond": record.source_fond,
        "source_inventory": record.source_inventory,
        "source_id": record.source_id,
        "language": record.language,
        "detected_language": record.detected_language,
        "language_source": record.language_source,
        "title_normalized": normalize_text(record.title_raw),
        "section_label": record.section_label,
        "chronology_basis": chronology_basis(record),
        "sheet_name": record.sheet_name,
        "excel_row": record.excel_row,
        "case_id_raw": record.case_id_raw,
        "case_id_normalized": record.case_id_normalized,
        "status": record.status,
        "section_year": record.section_year or "",
        "start_year": record.start_year or "",
        "end_year": record.end_year or "",
        "pages": record.pages or "",
        "title_raw": record.title_raw,
        "dates_raw": record.dates_raw,
        "pages_raw": record.pages_raw,
        "notes_raw": record.notes_raw,
        "category_ids": " | ".join(record.categories),
        "category_labels": " | ".join(record.category_labels),
        "macroblocks": " | ".join(record.macroblocks),
        "scores": scores_text,
        "evidence": evidence_text,
        "context_category_ids": " | ".join(record.context_categories),
        "context_category_labels": " | ".join(
            record.context_category_labels
        ),
        "context_evidence": context_evidence_text,
        "review_flags": " | ".join(record.review_flags),
    }


def chronology_basis(record: Record) -> str:
    if record.start_year is not None:
        if (not MIN_ALLOWED_YEAR <= record.start_year <= MAX_ALLOWED_YEAR
                or (record.end_year is not None and
                    (record.end_year < record.start_year or not MIN_ALLOWED_YEAR <= record.end_year <= MAX_ALLOWED_YEAR))):
            return "invalid_dates"
        return "dates_field"
    return "section_heading" if record.section_year else "unknown"


def get_decade(record: Record) -> int | None:
    if chronology_basis(record) == "invalid_dates":
        return None
    year = record.start_year or record.section_year
    return year // 10 * 10 if year is not None else None


def create_service_tables(records: list[Record], categories: list[Category]):
    record_fields = [
        "record_uid", "source_archive", "source_fond", "source_inventory", "source_id",
        "language", "detected_language", "language_source", "title_normalized",
        "section_label", "chronology_basis",
        "sheet_name", "excel_row", "case_id_raw", "case_id_normalized", "status",
        "section_year", "start_year", "end_year", "pages",
        "title_raw", "dates_raw", "pages_raw", "notes_raw",
        "category_ids", "category_labels", "macroblocks",
        "scores", "evidence", "context_category_ids",
        "context_category_labels", "context_evidence", "review_flags",
    ]
    active = [record for record in records if record.status == "case"]
    topic_unclassified = [record for record in active if not record.categories]
    context_only = [record for record in topic_unclassified if record.context_categories]
    unclassified = [record for record in topic_unclassified if not record.context_categories]
    needs_review = [record for record in active if record.review_flags]
    context_mentions = [record for record in active if record.context_categories]
    for filename, selected in (
        ("records.csv", records),
        ("topic_unclassified_cases.csv", topic_unclassified),
        ("unclassified_cases.csv", unclassified),
        ("context_only_cases.csv", context_only),
        ("needs_review.csv", needs_review),
        ("context_mentions.csv", context_mentions),
    ):
        written = write_csv(
            WORK_DIR / filename,
            record_fields,
            (record_to_row(record) for record in selected),
        )
        if written != len(selected):
            raise RuntimeError(
                f"{filename}: записано {written} рядків замість "
                f"{len(selected)}"
            )

    description_rows: list[dict[str, Any]] = []
    present_sheets = {record.sheet_name for record in records}
    ordered_sheets = [name for name in SHEET_NAMES if name in present_sheets]
    ordered_sheets += sorted(present_sheets.difference(ordered_sheets))
    for sheet_name in ordered_sheets:
        sheet_records = [record for record in records if record.sheet_name == sheet_name]
        sheet_active = [record for record in sheet_records if record.status == "case"]
        subject = sum(bool(record.categories) for record in sheet_active)
        context_only_count = sum(
            not record.categories and bool(record.context_categories)
            for record in sheet_active
        )
        description_rows.append(
            {
                "sheet_name": sheet_name,
                "description": configured_description_label(sheet_name),
                "cases_in_analysis": len(sheet_active),
                "withdrawn_records": sum(
                    record.status == "withdrawn" for record in sheet_records
                ),
                "subject_classified": subject,
                "context_only": context_only_count,
                "unclassified": len(sheet_active) - subject - context_only_count,
                "category_assignments": sum(
                    len(record.categories) for record in sheet_active
                ),
            }
        )
    write_csv(
        WORK_DIR / "description_summary.csv",
        [
            "sheet_name", "description", "cases_in_analysis",
            "withdrawn_records", "subject_classified", "context_only",
            "unclassified", "category_assignments",
        ],
        description_rows,
    )

    category_counts = Counter(
        category_id for record in active for category_id in record.categories
    )
    category_by_id = {category.id: category for category in categories}
    write_csv(
        WORK_DIR / "category_counts.csv",
        [
            "category_id", "category_label", "macroblock", "cases",
            "percent_of_analyzed_titles",
        ],
        (
            {
                "category_id": category.id,
                "category_label": category.label,
                "macroblock": category.macroblock,
                "cases": category_counts[category.id],
                "percent_of_analyzed_titles": (
                    f"{category_counts[category.id] / len(active) * 100:.2f}"
                    if active else "0.00"
                ),
            }
            for category in categories
        ),
    )

    context_counts = Counter(
        category_id
        for record in active
        for category_id in record.context_categories
    )
    write_csv(
        WORK_DIR / "context_category_counts.csv",
        [
            "category_id", "category_label", "cases",
            "percent_of_analyzed_titles",
        ],
        (
            {
                "category_id": category.id,
                "category_label": category.label,
                "cases": context_counts[category.id],
                "percent_of_analyzed_titles": (
                    f"{context_counts[category.id] / len(active) * 100:.2f}"
                    if active else "0.00"
                ),
            }
            for category in categories
        ),
    )

    decade_counts = Counter(
        decade for record in active
        if (decade := get_decade(record)) is not None
    )
    chronology_counts = Counter((get_decade(r), chronology_basis(r)) for r in active)
    write_csv(WORK_DIR / "chronology_sources.csv", ["decade", "basis", "titles"],
              ({"decade": decade if decade is not None else "", "basis": basis, "titles": count}
               for (decade, basis), count in sorted(chronology_counts.items(), key=lambda x: (x[0][0] or 0, x[0][1]))))
    write_csv(
        WORK_DIR / "cases_by_decade.csv",
        ["decade", "cases"],
        (
            {"decade": f"{decade}–{decade + 9}", "cases": count}
            for decade, count in sorted(decade_counts.items())
        ),
    )

    theme_decades: dict[str, Counter[int]] = defaultdict(Counter)
    for record in active:
        decade = get_decade(record)
        if decade is None:
            continue
        for category_id in record.categories:
            theme_decades[category_id][decade] += 1

    theme_rows: list[dict[str, Any]] = []
    for category_id, counts in theme_decades.items():
        for decade, count in sorted(counts.items()):
            theme_rows.append(
                {
                    "category_id": category_id,
                    "category_label": category_by_id[category_id].label,
                    "decade": f"{decade}–{decade + 9}",
                    "cases": count,
                }
            )
    write_csv(
        WORK_DIR / "themes_by_decade.csv",
        ["category_id", "category_label", "decade", "cases"],
        theme_rows,
    )
    return {
        "active": active,
        "topic_unclassified": topic_unclassified,
        "context_only": context_only,
        "unclassified": unclassified,
        "needs_review": needs_review,
        "context_mentions": context_mentions,
        "category_counts": category_counts,
        "context_counts": context_counts,
        "decade_counts": decade_counts,
        "theme_decades": theme_decades,
        "description_rows": description_rows,
    }


def one_line(value: str, limit: int = 500) -> str:
    value = re.sub(r"\s+", " ", value).strip()
    return value if len(value) <= limit else value[:limit - 1] + "…"


def write_error_log(issues: list[Issue], records: list[Record]) -> None:
    errors = [issue for issue in issues if issue.level == "ERROR"]
    warnings = [issue for issue in issues if issue.level == "WARNING"]
    statuses = Counter(record.status for record in records)
    lines = [
        "ПЕРЕВІРКА АРХІВНОГО ОПИСУ",
        f"Файл: {INPUT_FILE}",
        "Опрацьовані аркуші: "
        + ", ".join(sorted({record.sheet_name for record in records})),
        f"Дата запуску: {datetime.now():%d.%m.%Y %H:%M:%S}",
        "",
        "РЕЗЮМЕ",
        f"Справи, включені до аналізу: {statuses['case']}",
        f"Вибулі записи: {statuses['withdrawn']}",
        f"Структурні заголовки років: {statuses['year_heading']}",
        f"Загальні заголовки груп справ: {statuses['group_heading']}",
        f"Записи архівного опису: {statuses['inventory']}",
        f"Нерозпізнані рядки: {statuses['unknown']}",
        f"Помилки: {len(errors)}",
        f"Попередження: {len(warnings)}",
        "",
    ]

    number = 1
    for heading, group in (("ПОМИЛКИ", errors), ("ПОПЕРЕДЖЕННЯ", warnings)):
        lines.extend([heading, ""])
        if not group:
            lines.extend(["Не виявлено.", ""])
            continue
        for issue in group:
            location = []
            if issue.sheet_name:
                location.append(f"аркуш «{issue.sheet_name}»")
            if issue.row is not None:
                location.append(f"рядок {issue.row}")
            if issue.case_id:
                location.append(f"справа № {issue.case_id}")
            suffix = f" ({', '.join(location)})" if location else ""
            lines.append(f"{number}. {issue.message}{suffix}")
            if issue.field:
                lines.append(f"   Поле: {issue.field}")
            if issue.value:
                lines.append(f"   Значення: {one_line(issue.value)}")
            lines.append("")
            number += 1

    (REPORTS_DIR / "error.log").write_text(
        "\n".join(lines), encoding="utf-8-sig"
    )


def write_analysis_report(
    records: list[Record],
    issues: list[Issue],
    categories: list[Category],
    macroblock_labels: dict[str, str],
    table_data: dict[str, Any],
    figures_created: bool,
    dictionary_version: str,
) -> None:
    active = table_data["active"]
    unclassified = table_data["unclassified"]
    context_only = table_data["context_only"]
    needs_review = table_data["needs_review"]
    context_mentions = table_data["context_mentions"]
    category_counts = table_data["category_counts"]
    status_counts = Counter(record.status for record in records)
    classified = [record for record in active if record.categories]
    multilabel = [record for record in active if len(record.categories) > 1]
    years = [
        year for record in active
        for year in (record.start_year, record.end_year)
        if year is not None
    ]
    pages = [record.pages for record in active if record.pages is not None]
    error_count = sum(issue.level == "ERROR" for issue in issues)
    warning_count = sum(issue.level == "WARNING" for issue in issues)
    macroblock_counts = Counter(
        macroblock for record in active for macroblock in record.macroblocks
    )
    category_by_id = {category.id: category for category in categories}

    def percent(value: int, total: int) -> str:
        return f"{value / total * 100:.2f}%" if total else "0.00%"

    def format_int(value: int) -> str:
        return f"{value:,}".replace(",", " ")

    lines = [
        "# Результати аналізу архівних описів",
        "",
        f"Дата запуску: {datetime.now():%d.%m.%Y %H:%M:%S}.",
        f"Версія скрипта: `{SCRIPT_VERSION}`. Версія словників: "
        f"`{dictionary_version}`.",
        "Опрацьовані аркуші: "
        + ", ".join(sorted({record.sheet_name for record in records})) + ".",
        "",
        "## Склад масиву",
        "",
        "| Показник | Значення |",
        "|---|---:|",
        f"| Заголовки справ, включені до аналізу | {format_int(len(active))} |",
        f"| Вибулі записи | {format_int(status_counts['withdrawn'])} |",
        f"| Заголовки років | {format_int(status_counts['year_heading'])} |",
        f"| Загальні заголовки груп справ | "
        f"{format_int(status_counts['group_heading'])} |",
        f"| Записи про архівний опис | {format_int(status_counts['inventory'])} |",
        f"| Нерозпізнані рядки | {format_int(status_counts['unknown'])} |",
        f"| Помилки | {format_int(error_count)} |",
        f"| Попередження | {format_int(warning_count)} |",
        "",
        "## Розподіл за архівними описами",
        "",
        "| Архівний опис | Справ у аналізі | Вибулі | Предметну тему визначено | Лише контекст | Не визначено |",
        "|---|---:|---:|---:|---:|---:|",
    ]
    for row in table_data["description_rows"]:
        lines.append(
            f"| {row['description']} | {format_int(row['cases_in_analysis'])} | "
            f"{format_int(row['withdrawn_records'])} | "
            f"{format_int(row['subject_classified'])} | "
            f"{format_int(row['context_only'])} | "
            f"{format_int(row['unclassified'])} |"
        )

    lines.extend(
        [
        "",
        "## Хронологічні й кількісні показники",
        "",
        f"- Розпізнаний діапазон років: "
        f"{min(years) if years else '—'}–{max(years) if years else '—'}.",
        f"- Справ із розпізнаною кількістю аркушів: {format_int(len(pages))}.",
        f"- Сума аркушів лише за заповненими числовими значеннями: {format_int(sum(pages))}.",
        f"- Заголовків без заповненого поля дат: {sum(not r.dates_raw for r in active)}.",
        f"- Заголовків без розпізнаної кількості аркушів: {len(active)-len(pages)}.",
        "- Хронологія: початковий рік із поля дат; за його відсутності — рік групового заголовка. Це різні підстави, позначені в records.csv та chronology_sources.csv; рік групи не є крайньою датою справи. Зворотні діапазони та роки поза контрольними межами не включені до хронологічних графіків.",
        (
            f"- Середня кількість аркушів: {statistics.mean(pages):.2f}."
            if pages else "- Середня кількість аркушів: —."
        ),
        (
            f"- Медіанна кількість аркушів: {statistics.median(pages):.2f}."
            if pages else "- Медіанна кількість аркушів: —."
        ),
        "",
        "## Покриття тематичною класифікацією",
        "",
        f"- Предметну тематику визначено: {format_int(len(classified))} "
        f"({percent(len(classified), len(active))}).",
        f"- Визначено лише контекст осіб або установ: "
        f"{format_int(len(context_only))} "
        f"({percent(len(context_only), len(active))}).",
        f"- Не визначено ні тему, ні контекст: {format_int(len(unclassified))} "
        f"({percent(len(unclassified), len(active))}).",
        f"- Належить до кількох категорій: {format_int(len(multilabel))} "
        f"({percent(len(multilabel), len(active))}).",
        f"- Позначено для контекстної перевірки: {format_int(len(needs_review))}.",
        f"- Виявлено окремі контекстні згадки осіб або установ: "
        f"{format_int(len(context_mentions))} "
        f"({percent(len(context_mentions), len(active))}).",
        "",
        "## Тематичні категорії",
        "",
        "| Категорія | Справ | Частка справ, включених до аналізу |",
        "|---|---:|---:|",
        ]
    )
    for category_id, count in category_counts.most_common():
        lines.append(
            f"| {category_by_id[category_id].label} | "
            f"{format_int(count)} | {percent(count, len(active))} |"
        )

    lines.extend(
        [
            "",
            "Сума часток може перевищувати 100%, оскільки одна справа може належати "
            "до кількох категорій.",
            "",
            "## Макроблоки дисертації",
            "",
            "| Макроблок | Справ | Частка справ, включених до аналізу |",
            "|---|---:|---:|",
        ]
    )
    for macroblock, count in macroblock_counts.most_common():
        lines.append(
            f"| {macroblock_labels.get(macroblock, macroblock)} | "
            f"{format_int(count)} | {percent(count, len(active))} |"
        )

    lines.extend(
        [
            "",
            "## Створені матеріали",
            "",
            "- reports/error.log — перелік помилок і попереджень;",
            "- tables/records.csv — усі розпізнані рядки;",
            "- tables/description_summary.csv — контрольні підсумки за кожним "
            "архівним описом;",
            "- tables/unclassified_cases.csv — справи без теми та контексту;",
            "- tables/topic_unclassified_cases.csv — усі справи без предметної теми, включно з контекстними;",
            "- tables/context_only_cases.csv — справи, для яких визначено лише "
            "контекст осіб або установ;",
            "- tables/needs_review.csv — справи з неоднозначними термінами;",
            "- tables/context_mentions.csv — згадки осіб та установ, які самі "
            "по собі не визначають тему;",
            "- tables/category_counts.csv — кількість справ за категоріями;",
            "- tables/context_category_counts.csv — кількість контекстних "
            "згадок за категоріями;",
            "- tables/cases_by_decade.csv — хронологічний розподіл;",
            "- tables/themes_by_decade.csv — динаміка тем за десятиліттями;",
            "- thematic_exports/ — похідні XLSX/CSV-вибірки за всіма "
            "увімкненими категоріями, а також некласифіковані та службові записи.",
        ]
    )
    if figures_created:
        lines.append("- figures/ — графіки у форматах PNG і SVG.")
    else:
        lines.extend(
            [
                "",
                "> Графіки не створено, оскільки бібліотеку matplotlib "
                "не встановлено. Решта аналізу виконана.",
            ]
        )
    lines.extend(
        [
            "",
            "## Методичне застереження",
            "",
            f"Це автоматизована класифікація за словниками версії {dictionary_version}. "
            "Тематичні категорії описують предмет справи; назви установ, "
            "посади та станові означення, які лише називають учасника або "
            "автора документа, винесено в окремі контекстні мітки. "
            "Перед використанням числових результатів у дисертації потрібно "
            "перевірити needs_review.csv та unclassified_cases.csv, "
            "після чого скоригувати словники.",
            "",
            "Тематичні XLSX-файли є похідними дослідницькими вибірками, "
            "а не новими архівними описами. Первинним джерелом залишається "
            "input.xlsx. Одна справа може бути представлена у кількох файлах, "
            "оскільки класифікація є багатозначною.",
        ]
    )
    (REPORTS_DIR / "analysis_report.md").write_text(
        "\n".join(lines), encoding="utf-8"
    )


def save_figure(plt, name: str) -> None:
    plt.tight_layout()
    plt.savefig(FIGURES_DIR / f"{name}.png", dpi=220, bbox_inches="tight")
    plt.savefig(FIGURES_DIR / f"{name}.svg", bbox_inches="tight")
    plt.close()


CHART_PALETTES = (
    ("#4C78A8", "#6F98BC", "#91B5D0", "#B8D1E3", "#D8E8F5"),
    ("#F4E0BE", "#E6BE78", "#D39B48", "#B9782C", "#92591D"),
    ("#557A46", "#729661", "#91B181", "#B2CCAA", "#D6E5D1"),
    ("#70477C", "#8C6797", "#AA89B3", "#C8AECF", "#E3D5E7"),
)


def configured_description_label(sheet_name: str) -> str:
    """Return a complete archival reference for legends and control tables."""
    config = SOURCE_CONFIG.get(sheet_name, {})
    archive = clean_cell(config.get("archive")) or sheet_name
    fond = clean_cell(config.get("fond"))
    inventory = clean_cell(config.get("inventory"))
    reference = archive
    if fond:
        reference += f", ф. {fond}"
    if inventory:
        reference += f": опис {inventory}"
    elif sheet_name.lower().startswith("опис "):
        reference += f": {sheet_name.lower()}"
    return reference


def chart_sheet_order(records: Iterable[Record]) -> list[str]:
    present = {record.sheet_name for record in records}
    configured = [name for name in SHEET_NAMES if name in present]
    return configured + sorted(present.difference(configured))


def chart_sheet_colors(components: Iterable[str]) -> dict[str, str]:
    """Use one tonal family per archive while keeping descriptions distinct."""
    components = list(components)
    all_sheets = list(SHEET_NAMES) + sorted(
        set(components).difference(SHEET_NAMES)
    )
    source_order: list[str] = []
    source_sheets: dict[str, list[str]] = defaultdict(list)
    for sheet_name in all_sheets:
        source_id = SOURCE_CONFIG.get(sheet_name, {}).get("source_id", sheet_name)
        if source_id not in source_order:
            source_order.append(source_id)
        source_sheets[source_id].append(sheet_name)

    colors: dict[str, str] = {}
    for source_index, source_id in enumerate(source_order):
        palette = CHART_PALETTES[source_index % len(CHART_PALETTES)]
        sheets = source_sheets[source_id]
        if len(sheets) == 1:
            shades = [palette[2]]
        else:
            shades = [
                palette[round(index * (len(palette) - 1) / (len(sheets) - 1))]
                for index in range(len(sheets))
            ]
        colors.update(zip(sheets, shades))
    return {component: colors[component] for component in components}


def build_chart_breakdowns(records, active) -> dict[str, Any]:
    """Prepare mutually exclusive description segments for stacked charts."""
    components = chart_sheet_order(records)
    decades = {component: Counter() for component in components}
    categories = {component: Counter() for component in components}
    coverage = {component: Counter() for component in components}
    statuses = {component: Counter() for component in components}

    for record in active:
        decade = get_decade(record)
        if decade is not None:
            decades[record.sheet_name][decade] += 1
        categories[record.sheet_name].update(record.categories)
        if record.categories:
            coverage[record.sheet_name]["subject"] += 1
        elif record.context_categories:
            coverage[record.sheet_name]["context_only"] += 1
        else:
            coverage[record.sheet_name]["unclassified"] += 1

    for record in records:
        if record.status in {"case", "withdrawn"}:
            statuses[record.sheet_name][record.status] += 1

    return {
        "components": components,
        "colors": chart_sheet_colors(components),
        "labels": {
            component: configured_description_label(component)
            for component in components
        },
        "decades": decades,
        "categories": categories,
        "coverage": coverage,
        "statuses": statuses,
    }


def _label_color(hex_color: str) -> str:
    color = hex_color.lstrip("#")
    red, green, blue = (int(color[index:index + 2], 16) for index in (0, 2, 4))
    luminance = (0.299 * red + 0.587 * green + 0.114 * blue) / 255
    return "#202020" if luminance > 0.62 else "white"


def _segment_labels(
    values: list[int], totals: list[int], chart_maximum: int, horizontal: bool
) -> list[str]:
    minimum_scale_share = 0.018 if horizontal else 0.035
    return [
        str(value)
        if (
            value >= 10
            and value / max(total, 1) >= 0.025
            and value / max(chart_maximum, 1) >= minimum_scale_share
        )
        else ""
        for value, total in zip(values, totals)
    ]


def draw_stacked_bars(
    ax,
    labels: list[str],
    components: list[str],
    values: dict[str, list[int]],
    colors: dict[str, str],
    component_labels: dict[str, str],
    *,
    horizontal: bool = False,
) -> list[int]:
    totals = [
        sum(values[component][index] for component in components)
        for index in range(len(labels))
    ]
    chart_maximum = max(totals, default=0)
    offsets = [0] * len(labels)
    for component in components:
        segment = values[component]
        kwargs = {
            "label": component_labels[component],
            "color": colors[component],
            "edgecolor": "white",
            "linewidth": 0.2,
        }
        if horizontal:
            bars = ax.barh(labels, segment, left=offsets, **kwargs)
        else:
            bars = ax.bar(labels, segment, bottom=offsets, **kwargs)
        ax.bar_label(
            bars,
            labels=_segment_labels(
                segment, totals, chart_maximum, horizontal
            ),
            label_type="center",
            fontsize=7,
            color=_label_color(colors[component]),
        )
        offsets = [offset + value for offset, value in zip(offsets, segment)]

    margin = max(chart_maximum * 0.012, 0.5)
    for index, total in enumerate(totals):
        if not total:
            continue
        total_label = f"Усього: {total}"
        if horizontal:
            ax.text(
                total + margin, index, total_label,
                va="center", fontsize=8, fontweight="bold",
            )
        else:
            ax.text(
                index, total + margin, total_label,
                ha="center", fontsize=8, fontweight="bold",
            )
    ax.legend(
        title="Архівний опис",
        loc="lower center",
        bbox_to_anchor=(0.5, 1.03),
        ncol=min(len(components), 3),
        frameon=False,
        fontsize=8,
        title_fontsize=8,
    )
    return totals


def create_figures(plt, records, categories, table_data) -> bool:
    if plt is None:
        print(
            "Увага: matplotlib не встановлено, тому графіки пропущено.\n"
            f'Для графіків виконайте: "{sys.executable}" -m pip install matplotlib'
        )
        return False

    plt.rcParams.update(
        {
            "font.family": "DejaVu Sans",
            "axes.titlesize": 13,
            "axes.labelsize": 10,
            "figure.dpi": 120,
        }
    )
    active = table_data["active"]
    category_counts = table_data["category_counts"]
    decade_counts = table_data["decade_counts"]
    theme_decades = table_data["theme_decades"]
    category_by_id = {category.id: category for category in categories}
    breakdowns = build_chart_breakdowns(records, active)
    components = breakdowns["components"]
    colors = breakdowns["colors"]
    component_labels = breakdowns["labels"]

    decades = sorted(decade_counts)
    decade_labels = [f"{decade}–{decade + 9}" for decade in decades]
    decade_values = {
        component: [breakdowns["decades"][component][decade] for decade in decades]
        for component in components
    }
    _, ax = plt.subplots(figsize=(10, 5.6))
    draw_stacked_bars(
        ax, decade_labels, components, decade_values, colors, component_labels
    )
    ax.set_title("Розподіл заголовків за десятиліттями")
    ax.set_xlabel("Початковий рік із дат або, за його відсутності, рік групи")
    ax.set_ylabel("Кількість справ")
    ax.tick_params(axis="x", rotation=45)
    for label in ax.get_xticklabels():
        label.set_ha("right")
    ax.grid(axis="y", alpha=0.25)
    ax.set_axisbelow(True)
    save_figure(plt, "cases_by_decade")

    ordered = category_counts.most_common()
    category_ids = [item[0] for item in ordered][::-1]
    labels = [category_by_id[category_id].label for category_id in category_ids]
    category_values = {
        component: [
            breakdowns["categories"][component][category_id]
            for category_id in category_ids
        ]
        for component in components
    }
    _, ax = plt.subplots(figsize=(10, 6.5))
    draw_stacked_bars(
        ax, labels, components, category_values, colors, component_labels,
        horizontal=True,
    )
    ax.set_title("Тематична структура заголовків справ")
    ax.set_xlabel("Кількість справ")
    ax.grid(axis="x", alpha=0.25)
    ax.set_axisbelow(True)
    save_figure(plt, "thematic_categories")

    coverage_keys = ["subject", "context_only", "unclassified"]
    coverage_labels = [
        "Предметну тему\nвизначено", "Лише контекст", "Не визначено"
    ]
    coverage_values = {
        component: [
            breakdowns["coverage"][component][key] for key in coverage_keys
        ]
        for component in components
    }
    _, ax = plt.subplots(figsize=(8.4, 5.2))
    draw_stacked_bars(
        ax, coverage_labels, components, coverage_values, colors,
        component_labels,
    )
    ax.set_title("Покриття класифікаційними словниками")
    ax.set_ylabel("Кількість справ")
    ax.grid(axis="y", alpha=0.25)
    ax.set_axisbelow(True)
    save_figure(plt, "classification_coverage")

    status_keys = ["case", "withdrawn"]
    status_labels = ["Включено до аналізу", "Вибулі"]
    status_values = {
        component: [
            breakdowns["statuses"][component][key] for key in status_keys
        ]
        for component in components
    }
    _, ax = plt.subplots(figsize=(8.4, 5.2))
    draw_stacked_bars(
        ax, status_labels, components, status_values, colors, component_labels
    )
    ax.set_title("Справи, включені до аналізу, та вибулі записи")
    ax.set_ylabel("Кількість записів")
    ax.grid(axis="y", alpha=0.25)
    ax.set_axisbelow(True)
    save_figure(plt, "included_and_withdrawn")

    if decades and ordered:
        category_ids = [category_id for category_id, _ in ordered]
        matrix = [
            [theme_decades[category_id][decade] for decade in decades]
            for category_id in category_ids
        ]
        plt.figure(figsize=(11, 7))
        image = plt.imshow(matrix, aspect="auto", cmap="Blues")
        plt.colorbar(image, label="Кількість справ")
        plt.xticks(
            range(len(decades)),
            [f"{decade}–{str(decade + 9)[-2:]}" for decade in decades],
            rotation=45, ha="right",
        )
        plt.yticks(
            range(len(category_ids)),
            [category_by_id[category_id].label for category_id in category_ids],
        )
        plt.title("Динаміка тематичних категорій за десятиліттями")
        plt.xlabel("Десятиліття")
        save_figure(plt, "themes_by_decade")
    return True


def print_summary(records, issues, table_data, figures_created) -> None:
    statuses = Counter(record.status for record in records)
    errors = sum(issue.level == "ERROR" for issue in issues)
    warnings = sum(issue.level == "WARNING" for issue in issues)
    active = table_data["active"]
    classified = sum(bool(record.categories) for record in active)
    print("\nАНАЛІЗ ЗАВЕРШЕНО")
    print(f"Справи в аналізі:        {statuses['case']}")
    print(f"Вибулі записи:           {statuses['withdrawn']}")
    print(f"Предметну тему визначено:{classified:>7}")
    print(f"Лише контекст:           {len(table_data['context_only']):>7}")
    print(f"Не визначено:            {len(table_data['unclassified']):>7}")
    print(f"Потребують перегляду:    {len(table_data['needs_review'])}")
    print(f"Контекстні згадки:       {len(table_data['context_mentions'])}")
    print(f"Помилки:                 {errors}")
    print(f"Попередження:            {warnings}")
    print(f"Графіки створено:        {'так' if figures_created else 'ні'}")
    print("\nОсновні результати:")
    print(f"  {REPORTS_DIR / 'error.log'}")
    print(f"  {REPORTS_DIR / 'analysis_report.md'}")
    print(f"  {WORK_DIR / 'records.csv'}")
    print(f"  {RUN_DIR / 'tables' / 'classification_index.csv'}")
    print(f"  {THEMATIC_DIR}")
    if figures_created:
        print(f"  {FIGURES_DIR}")


def main() -> None:
    global RUN_DIR, REPORTS_DIR, FIGURES_DIR, WORK_DIR, THEMATIC_DIR
    run_id = datetime.now().strftime("%Y%m%d_%H%M%S_%f")
    print(
        f"=== АНАЛІЗ АРХІВНИХ ОПИСІВ "
        f"— {SCRIPT_VERSION} ==="
    )
    print(f"Папка скрипта: {BASE_DIR}")
    openpyxl_module, yaml_module, plt = load_dependencies()
    configure_analysis(yaml_module)
    configure_sources(yaml_module)
    print(f"Вхідний файл:  {INPUT_FILE}")
    print(f"Аркуші:        {', '.join(SHEET_NAMES)}")
    RUN_DIR = OUTPUTS_DIR / run_id
    REPORTS_DIR = RUN_DIR / "reports"
    FIGURES_DIR = RUN_DIR / "figures"
    WORK_DIR = RUN_DIR / "tables"
    THEMATIC_DIR = RUN_DIR / "thematic_exports"
    ensure_directories()

    print("\n1. Завантаження класифікаційних словників...")
    (
        categories, ambiguities, macroblock_labels,
        dictionary_version,
    ) = load_dictionaries(yaml_module)
    LANGUAGE_CATEGORIES.clear()
    LANGUAGE_AMBIGUITIES.clear()
    LANGUAGE_CATEGORIES["uk"] = categories
    LANGUAGE_AMBIGUITIES["uk"] = ambiguities
    ru_categories, ru_ambiguities, _, _ = load_dictionaries(yaml_module, "ru")
    LANGUAGE_CATEGORIES["ru"] = ru_categories
    LANGUAGE_AMBIGUITIES["ru"] = ru_ambiguities
    print(f"   Завантажено категорій: {len(categories)}")

    print("2. Читання та перевірка input.xlsx...")
    records, issues, _headers = read_records(openpyxl_module)
    for source in sorted({r.source_id for r in records}):
        subset = [r for r in records if r.source_id == source and r.status == "case"]
        missing_pages = sum(not r.pages_raw for r in subset)
        if subset and missing_pages and not SOURCE_CONFIG[subset[0].sheet_name].get("pages_expected", True):
            issues.append(Issue("WARNING", None, "", "Кількість аркушів", str(missing_pages),
                                f"{source}: кількість аркушів не заповнено у {missing_pages} заголовках; зведене попередження.", subset[0].sheet_name))
    print(f"   Розпізнано рядків: {len(records)}")

    print("3. Тематична класифікація...")
    macroblock_order = list(macroblock_labels)
    classify_records(records, categories, ambiguities, macroblock_order)

    roots = REPORTS_DIR, FIGURES_DIR, WORK_DIR
    scopes = {source: [r for r in records if r.source_id == source]
              for source in sorted({r.source_id for r in records})}
    scopes["combined"] = records
    summary_rows = []
    for scope, subset in scopes.items():
        print(f"4–7. Таблиці, графіки та звіт: {scope}...", flush=True)
        REPORTS_DIR, FIGURES_DIR, WORK_DIR = [p / scope for p in roots]
        ensure_directories()
        sheets = {r.sheet_name for r in subset}
        scoped_issues = [i for i in issues if i.sheet_name in sheets or not i.sheet_name] if scope != "combined" else issues
        table_data = create_service_tables(subset, categories)
        write_error_log(scoped_issues, subset)
        figures_created = create_figures(plt, subset, categories, table_data)
        write_analysis_report(subset, scoped_issues, categories, macroblock_labels,
                              table_data, figures_created, dictionary_version)
        cases = table_data["active"]
        summary_rows.append({"source": scope, "titles": len(cases),
                             "subject": sum(bool(r.categories) for r in cases),
                             "context_only": len(table_data["context_only"]),
                             "no_evidence": len(table_data["unclassified"]),
                             "missing_dates": sum(not r.dates_raw for r in cases),
                             "missing_pages": sum(r.pages is None for r in cases)})
    write_csv(roots[2] / "source_summary.csv", list(summary_rows[0]), summary_rows)
    print("8. Формування тематичних XLSX/CSV-вибірок...", flush=True)
    write_combined_indexes(records, categories, roots[2])
    export_manifest = export_thematic_selections(
        openpyxl_module, records, categories
    )

    tracked = (
        [INPUT_FILE, Path(__file__), BASE_DIR / "models.py", BASE_DIR / "text_matching.py"]
        + sorted(CONFIG_DIR.glob("*.yaml"))
        + sorted(DICTIONARIES_DIR.rglob("*.yaml"))
    )
    manifest = {
        "status": "complete",
        "script_version": SCRIPT_VERSION,
        "dictionary_version": dictionary_version,
        "run_id": run_id,
        "completed_at": datetime.now().isoformat(timespec="seconds"),
        "python": platform.python_version(),
        "openpyxl": openpyxl_module.__version__,
        "pyyaml": yaml_module.__version__,
        "configured_sheets": list(SHEET_NAMES),
        "enabled_categories": [category.id for category in categories],
        "source_summary": summary_rows,
        "thematic_exports": export_manifest,
        "sha256": {
            str(path.relative_to(BASE_DIR)): hashlib.sha256(
                path.read_bytes()
            ).hexdigest()
            for path in tracked
        },
    }
    (RUN_DIR / "run_manifest.json").write_text(
        json.dumps(manifest, ensure_ascii=False, indent=2), encoding="utf-8"
    )
    print_summary(records, issues, table_data, figures_created)
    print(f"Ідентифікатор запуску: {run_id}. Старі результати збережено.")
    print(f"Повна папка запуску:    {RUN_DIR}")


if __name__ == "__main__":
    try:
        main()
    except SystemExit:
        raise
    except Exception as error:
        print("\nКРИТИЧНА ПОМИЛКА")
        print(str(error))
        print("\nТехнічні подробиці:")
        traceback.print_exc()
        raise SystemExit(1)
